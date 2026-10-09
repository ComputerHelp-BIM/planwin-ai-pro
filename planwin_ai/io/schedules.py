"""Drawing-office member schedules (Excel): columns, beams, footings, slabs and walls.

One sheet per schedule; a sheet is left out when the model has no such members.

* **Column schedule** – one row per group of column marks that are identical at every
  level, one column per level range over which no row changes (the grouping of the
  detail drawings, :func:`planwin_ai.io.detail_dxf.column_schedule`).  Each cell reads
  ``300×450 / 8-16Ø / 8Ø@150 (conf. 100, l0 450)``.  **Column list** has one row per
  designed column segment.
* **Beam schedule** – per level, physical beams (all spans of one plan beam) with the
  same size, spans (within 50 mm) and bars share a row group, e.g. "B1, B4, B7"; a beam
  listed reversed (top bars swapped end for end) counts as identical.  One row per span.
* **Footing schedule** – types F1, F2 … of identical isolated footings
  (:func:`planwin_ai.io.detail_dxf.footing_types`) with the columns served and the
  bearing pressure check; **Combined footings** when there are any.
* **Slab schedule** and **Wall schedule** – one row per panel / wall storey.

Forces and pressures are shown in the display units (:data:`planwin_ai.units.current`).
"""

from __future__ import annotations

from ..core.model import Project
from ..design.report import BeamDesign, DesignReport
from .boq_excel import footnote, header_block, page_setup, table_header
from .detail_dxf import ColumnCell, column_schedule, footing_types, levels_label, natural_key

_SLAB_KIND = {"two_way": "Two-way", "one_way": "One-way", "cantilever": "Cantilever"}
NOTE = "Read with the structural drawings and general notes. Bar marks: n-16Ø = n bars of 16 mm; 8Ø@150 = links."


def _mm(v: float) -> str:
    return f"{v * 1000:.0f}"


def column_cell_text(cell: ColumnCell | None) -> str:
    """``300×450 / 8-16Ø / 8Ø@150 (conf. 100, l0 450)`` – size, main bars, ties and confining hoops."""
    if cell is None:
        return "–"
    out = f"{_mm(cell.b)}×{_mm(cell.d)} / {cell.bars.count}-{cell.bars.dia}Ø"
    t, c = cell.tie, cell.tie_confined
    out += f" / {t.dia}Ø@{t.spacing:.0f}" if t else " / ties: revise"
    if c is not None and cell.l0 > 0:
        sp = f"{c.spacing:.0f}" if t and c.dia == t.dia else f"{c.dia}Ø@{c.spacing:.0f}"
        out += f" (conf. {sp}, l0 {_mm(cell.l0)})"
    return out


def _beam_sig(segs: list[BeamDesign], reverse: bool = False) -> tuple:
    seq = segs[::-1] if reverse else segs
    return tuple(
        (
            round(b.b, 3),
            round(b.d, 3),
            round(b.span / 0.05),  # spans within 50 mm
            b.bottom,
            *((b.top_r, b.top_l) if reverse else (b.top_l, b.top_r)),
            str(b.links or b.stirrups),
            str(b.links_end or ""),
            b.side_face,
        )
        for b in seq
    )


def beam_schedule_groups(rep: DesignReport) -> list[tuple[int, str, list[str], list[BeamDesign]]]:
    """(level index, level name, marks, spans of the representative beam) per group of identical beams."""
    phys: dict[tuple[int, str], list[BeamDesign]] = {}
    for b in rep.beams:
        phys.setdefault((b.level_index, b.group or b.mark), []).append(b)
    groups: dict[tuple, list[list[BeamDesign]]] = {}
    for (li, _), segs in phys.items():
        segs.sort(key=lambda b: b.member_id)  # segments are numbered along the beam
        key = (li, min(_beam_sig(segs), _beam_sig(segs, True)))
        groups.setdefault(key, []).append(segs)
    out = []
    for (li, _), beams in groups.items():
        marks = sorted(dict.fromkeys(s[0].mark for s in beams), key=natural_key)
        rep_segs = min(beams, key=lambda s: natural_key(s[0].mark))
        out.append((li, rep_segs[0].level, marks, rep_segs))
    out.sort(key=lambda g: (g[0], natural_key(g[2][0])))
    return out


def write_schedules_excel(path: str, project: Project, rep: DesignReport, watermark: str = "") -> str:
    """Write the member schedules workbook (see module docs) and return ``path``."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side

    from .. import units as _units

    u = _units.current
    F, A = u.label("force"), u.label("area")
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="A6A6A6")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(title: str, headers: list[str], rows: list[list], widths: list[float], note: str = "") -> None:
        if not rows:
            return
        ws = wb.create_sheet(title)
        hrow = header_block(ws, f"{title} – {project.name}", project, watermark)
        table_header(ws, hrow, headers)
        for r in rows:
            ws.append(r)
            for c in ws[ws.max_row]:
                c.border, c.alignment = box, wrap
                if isinstance(c.value, float):
                    c.number_format = "0.00"
        footnote(ws, *(n for n in (note, NOTE) if n))
        page_setup(ws, hrow, widths, landscape=True)

    # ---- columns
    ranges, crow = column_schedule(project, rep)
    heads = [levels_label(project, range(a, b + 1)) for a, b in ranges]
    sheet(
        "Column schedule",
        ["Column marks", *heads],
        [[", ".join(r.marks), *(column_cell_text(c) for c in r.cells)] for r in crow],
        [22] + [30] * len(heads),
        "Cell: b×D (mm) / main bars / ties @ spacing (confining hoops within l0 at each end, IS 13920 cl 8).",
    )
    cols = sorted(rep.columns, key=lambda c: (c.level_index, natural_key(c.mark)))
    sheet(
        "Column list",
        ["Level", "Mark", "b mm", "D mm", "Main bars", "Ties", "Confinement (l0)", "l0 m"]
        + [f"Pu {F}", "Utilisation", "OK"],
        [
            [
                c.level,
                c.mark,
                round(c.b * 1000),
                round(c.d * 1000),
                c.bars,
                str(c.tie) if c.tie else c.ties,
                (c.ties or "").split(" / ")[0] if c.tie_confined else "",
                round(c.l0, 3) if c.tie_confined else "",
                round(u.show(c.Pu, "force"), 1),
                round(c.utilisation, 3),
                "YES" if c.ok else "NO",
            ]
            for c in cols
        ],
        [12, 8, 7, 7, 18, 18, 34, 7, 11, 11, 6],
    )
    # ---- beams
    rows = []
    for _li, level, marks, segs in beam_schedule_groups(rep):
        n = len(segs)
        for j, b in enumerate(segs):
            remarks = "; ".join(dict.fromkeys(nt for s in segs for nt in s.notes)) if j == 0 else ""
            if not b.ok:
                remarks = ("REVISE – " + remarks).strip(" –")
            rows.append(
                [
                    level if j == 0 else "",
                    ", ".join(marks) if j == 0 else "",
                    f"{_mm(b.b)}×{_mm(b.d)}",
                    f"{j + 1} of {n}" if n > 1 else "1",
                    round(b.span, 3),
                    b.bottom,
                    b.top_l,
                    b.top_r,
                    str(b.links_end) if b.links_end else "as mid-span",
                    str(b.links) if b.links else b.stirrups,
                    b.side_face or "–",
                    remarks,
                ]
            )
    sheet(
        "Beam schedule",
        ["Level", "Marks", "b×D mm", "Span", "Span m", "Bottom", "Top left", "Top right"]
        + ["Stirrups end zone (2d)", "Stirrups mid-span", "Side face / torsion", "Remarks"],
        rows,
        [11, 22, 10, 7, 8, 10, 10, 10, 20, 20, 24, 30],
        "Spans are numbered from the beam's first end; beams listed reversed have their top bars swapped.",
    )
    # ---- footings
    sbc = project.design.sbc
    by_mark = {f.mark: f for f in rep.footings}
    rows = []
    for t in footing_types(rep):
        fs = [by_mark[m] for m in t.columns if m in by_mark]
        q = max((f.q for f in fs), default=0.0)
        notes = "; ".join(dict.fromkeys(n for f in fs for n in f.notes))
        rows.append(
            [
                t.mark,
                ", ".join(t.columns),
                len(t.columns),
                f"{_mm(t.L)}×{_mm(t.B)}×{_mm(t.D)}",
                str(t.mesh_L) if t.mesh_L else fs[0].bars_L,
                str(t.mesh_B) if t.mesh_B else fs[0].bars_B,
                round(u.show(max(f.P_service for f in fs), "force"), 1),
                round(u.show(q, "area"), 1),
                round(u.show(sbc, "area"), 1),
                "OK" if q <= sbc * 1.001 else "EXCEEDS SBC",
                ("REVISE – " if not all(f.ok for f in fs) else "") + notes,
            ]
        )
    sheet(
        "Footing schedule",
        ["Type", "Columns served", "No.", "L×B×D mm", "Bars along L (bottom)", "Bars along B (bottom)"]
        + [f"P service max {F}", f"q max {A}", f"SBC {A}", "SBC check", "Remarks"],
        rows,
        [7, 30, 6, 16, 20, 20, 13, 11, 10, 12, 30],
        "L is along the column depth D. Identical isolated footings share a type.",
    )
    sheet(
        "Combined footings",
        ["Columns", "L×B×D mm", f"P service {F}", f"q {A}", "SBC check", "Top (between columns)", "Bottom"]
        + ["Transverse", "OK", "Remarks"],
        [
            [
                " + ".join(c.marks),
                f"{_mm(c.L)}×{_mm(c.B)}×{_mm(c.D)}",
                round(u.show(c.P_service, "force"), 1),
                round(u.show(c.q, "area"), 1),
                "OK" if c.q <= sbc * 1.001 else "EXCEEDS SBC",
                c.top,
                c.bottom,
                "; ".join(c.transverse),
                "YES" if c.ok else "NO",
                "; ".join(c.notes),
            ]
            for c in rep.combined_footings
        ],
        [16, 18, 12, 10, 12, 24, 24, 30, 6, 30],
    )
    # ---- slabs
    plan_order = {p.name: i for i, p in enumerate(project.plans)}
    slabs = sorted(rep.slabs, key=lambda ps: (plan_order.get(ps[0], 99), natural_key(ps[1].mark)))
    sheet(
        "Slab schedule",
        ["Plan", "Mark", "lx×ly m", "Type", "D mm", "Short span (bottom)", "Long span (bottom)", "Top (supports)"]
        + ["OK", "Remarks"],
        [
            [
                pn,
                s.mark,
                f"{s.lx:.2f}×{s.ly:.2f}",
                _SLAB_KIND.get(s.kind, s.kind),
                round(s.D_mm),
                s.ast_x,
                s.ast_y,
                s.ast_neg or "–",
                "YES" if s.ok else "NO",
                "; ".join(s.notes),
            ]
            for pn, s in slabs
        ],
        [12, 8, 12, 11, 7, 20, 20, 20, 6, 30],
    )
    # ---- walls
    walls = sorted(rep.walls, key=lambda w: (natural_key(w.mark), w.level_index))
    sheet(
        "Wall schedule",
        ["Mark", "Level", "t×L mm", "Height m", "Curtains", "Vertical", "Horizontal", "Boundary elements"]
        + ["Utilisation", "OK", "Remarks"],
        [
            [
                w.mark,
                w.level,
                f"{_mm(w.t)}×{_mm(w.Lw)}",
                round(w.height, 3),
                w.curtains,
                w.vertical,
                w.horizontal,
                w.boundary,
                round(w.utilisation, 3),
                "YES" if w.ok else "NO",
                "; ".join(w.notes),
            ]
            for w in walls
        ],
        [8, 11, 12, 9, 9, 22, 22, 48, 11, 6, 30],
    )
    if not wb.sheetnames:
        raise ValueError("nothing to schedule – run the design first")
    wb.save(path)
    return path
