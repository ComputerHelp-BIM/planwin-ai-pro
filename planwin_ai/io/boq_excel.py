"""Standalone, editable BOQ & cost estimate workbook (Excel) with live formulas.

Sheets
------
* **Rates** – every rate of ``project.design.rates`` (plus PCC and any concrete grade the
  quantities use) as yellow input cells.  Everything else prices from these cells.
* **BOQ** – Indian BOQ layout (Item No., Description, Unit, Quantity, Rate, Amount) grouped
  by member type: RCC per grade, reinforcement, centering & shuttering (and PCC under the
  footings).  Rate = ``=Rates!$D$n``, Amount = ``=Qty*Rate``, ``SUBTOTAL`` per group and a
  grand total, plus total quantities and the steel / RCC ratio as formulas.
* **By floor** / **By type** – concrete per grade, steel and formwork per level / member
  type; RCC total, kg/m³ and cost are formulas on the Rates sheet.
* **Revisions** – saved BOQ revisions (``project.meta["revisions"]``) and, with two or
  more, the last two compared with change and change % formulas.

Quantities are those of ``rep.boq`` (:func:`planwin_ai.design.quantities.quantities`), so the
grand total equals ``rep.boq["cost"]`` at the project rates.  The header helpers here are
shared with :mod:`planwin_ai.io.schedules`.
"""

from __future__ import annotations

import datetime as _dt

from .. import APP_NAME, __version__
from ..core.model import Project
from ..design.quantities import _Ledger, compare_revisions
from ..design.report import DesignReport
from .report_common import DISCLAIMER

NAVY = "1F3A5F"
INPUT_FILL = "FFF2CC"  # light yellow: cells meant to be edited
SUB_FILL = "DCE6F1"
F_M3 = "#,##0.00"
F_KG = "#,##0"
F_M2 = "#,##0.00"
F_INR = '"₹" #,##0'
F_RATE = '"₹" #,##0.00'
F_PCT = "0.0%"
UNIT_FMT = {"m³": F_M3, "kg": F_KG, "m²": F_M2, "₹": F_INR}
_KIND = {"Columns": "columns", "Walls": "shear walls", "Beams": "beams", "Slabs": "slabs", "Footings": "footings"}


# ----------------------------------------------------------------------------- shared sheet helpers
def header_block(ws, title: str, project: Project, watermark: str = "") -> int:
    """Title and project block at the top of ``ws``; returns the row for the table header."""
    from openpyxl.styles import Font

    ws.append([title])
    ws.cell(1, 1).font = Font(size=14, bold=True, color=NAVY)
    info = [
        ("Project", project.name),
        ("Client", project.client),
        ("Engineer", project.engineer),
        ("Location", project.location),
        ("Date", _dt.date.today().isoformat()),
        ("Prepared with", f"{APP_NAME} {__version__}"),
    ]
    if watermark:
        info.append(("Licence", watermark))
    for k, v in info:
        if k in ("Client", "Engineer", "Location") and not v:
            continue
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    return ws.max_row + 2  # one blank row before the table


def table_header(ws, row: int, headers: list[str]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill, font = PatternFill("solid", fgColor=NAVY), Font(color="FFFFFF", bold=True)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h)
        c.fill, c.font = fill, font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def page_setup(ws, header_row: int, widths: list[float], landscape: bool = False) -> None:
    """Column widths, frozen header, A4 fit-to-width print with the header row repeated."""
    from openpyxl.utils import get_column_letter

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row + 1}"
    ps = ws.page_setup
    ps.paperSize = ws.PAPERSIZE_A4
    ps.orientation = "landscape" if landscape else "portrait"
    ps.fitToWidth, ps.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Page &P of &N"


def footnote(ws, *lines: str) -> None:
    from openpyxl.styles import Font

    ws.append([])
    for ln in (*lines, DISCLAIMER):
        ws.append([ln])
        ws.cell(ws.max_row, 1).font = Font(italic=True, color="595959")


def bold_row(ws, row: int, fill: str | None = None) -> None:
    from openpyxl.styles import Font, PatternFill

    for c in ws[row]:
        c.font = Font(bold=True)
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)


# ----------------------------------------------------------------------------- BOQ
def rate_key(grade: str) -> str:
    """Key of ``project.design.rates`` that prices concrete ``grade`` (as the quantity ledger does)."""
    return "concrete_pcc" if grade.startswith("PCC") else f"concrete_{grade.lower().replace(' ', '_')}"


def _rate_label(key: str, fy: float) -> tuple[str, str]:
    if key == "concrete_pcc":
        return "PCC M10 (1:3:6) concrete", "m³"
    if key.startswith("concrete_"):
        return f"RCC {key[9:].upper().replace('_', ' ')} concrete (design mix)", "m³"
    if key == "steel_kg":
        return f"Reinforcement steel Fe{fy:.0f} (TMT), cut, bent and placed", "kg"
    if key == "formwork_m2":
        return "Centering & shuttering (formwork)", "m²"
    return key.replace("_", " "), ""


def _rates_sheet(wb, project: Project, boq: dict, watermark: str) -> dict[str, str]:
    """Rates sheet; returns {rate key: absolute reference} for every rate the BOQ uses."""
    from openpyxl.styles import PatternFill

    rates = project.design.rates
    led = _Ledger(rates)
    fy = project.design.fy_main
    rows: dict[str, tuple[float, str]] = {k: (float(v), "") for k, v in rates.items()}
    for grade in boq.get("concrete", {}):  # grades priced through a fallback rate get their own input
        k = rate_key(grade)
        if k not in rows:
            note = "" if k == "concrete_pcc" else f"not in the project rates – {led.rate(grade):,.0f} used"
            rows[k] = (led.rate(grade), note)
    rows.setdefault("steel_kg", (rates.get("steel_kg", 75.0), "default"))
    rows.setdefault("formwork_m2", (rates.get("formwork_m2", 550.0), "default"))
    ws = wb.create_sheet("Rates")
    hrow = header_block(ws, "Schedule of rates (inputs)", project, watermark)
    table_header(ws, hrow, ["Code", "Description", "Unit", "Rate (₹)", "Note"])
    refs: dict[str, str] = {}
    fill = PatternFill("solid", fgColor=INPUT_FILL)
    for k, (v, note) in rows.items():
        desc, unit = _rate_label(k, fy)
        ws.append([k, desc, unit, v, note])
        c = ws.cell(ws.max_row, 4)
        c.fill, c.number_format = fill, F_RATE
        refs[k] = f"Rates!$D${ws.max_row}"
    footnote(ws, "Yellow cells are inputs: change a rate and every amount in the workbook updates.")
    page_setup(ws, hrow, [18, 52, 8, 14, 40])
    return refs


def _ordered_grades(by_grade: dict) -> list[str]:
    return sorted(by_grade, key=lambda g: (g.startswith("PCC"), g))


def _boq_sheet(wb, project: Project, boq: dict, refs: dict[str, str], watermark: str) -> None:
    ws = wb.create_sheet("BOQ", 0)
    hrow = header_block(ws, "Bill of quantities & cost estimate", project, watermark)
    table_header(ws, hrow, ["Item No.", "Description", "Unit", "Quantity", "Rate (₹)", "Amount (₹)"])
    fy = project.design.fy_main
    first = hrow + 1
    cells: dict[str, list[str]] = {"rcc": [], "pcc": [], "steel": [], "formwork": []}
    letters = iter("ABCDEFGHIJ")
    for kind, v in (boq.get("by_type") or {}).items():
        if not (v.get("by_grade") or v.get("steel") or v.get("formwork")):
            continue  # e.g. no walls
        letter, what = next(letters), _KIND.get(kind, kind.lower())
        ws.append([letter, kind.upper()])
        bold_row(ws, ws.max_row)
        items: list[tuple[str, str, float, str, str]] = []
        for g in _ordered_grades(v.get("by_grade", {})):
            if g.startswith("PCC"):
                desc = f"{g} (1:3:6) levelling course below {what}, 100 mm thick"
                items.append((desc, "m³", v["by_grade"][g], rate_key(g), "pcc"))
            else:
                desc = f"RCC {g} in {what} (formwork and reinforcement measured separately)"
                items.append((desc, "m³", v["by_grade"][g], rate_key(g), "rcc"))
        laps = {"Columns": ", incl. 10 % laps", "Walls": ", incl. 10 % laps", "Slabs": ", incl. laps & chairs"}
        desc = f"Reinforcement steel Fe{fy:.0f} in {what}{laps.get(kind, '')}"
        items.append((desc, "kg", v["steel"], "steel_kg", "steel"))
        items.append((f"Centering & shuttering – {what}", "m²", v["formwork"], "formwork_m2", "formwork"))
        top = ws.max_row + 1
        for i, (desc, unit, qty, key, head) in enumerate(items, 1):
            r = ws.max_row + 1
            ws.append([f"{letter}.{i}", desc, unit, round(qty, 3), f"={refs[key]}", f"=D{r}*E{r}"])
            ws.cell(r, 4).number_format = UNIT_FMT[unit]
            ws.cell(r, 5).number_format = F_RATE
            ws.cell(r, 6).number_format = F_INR
            cells[head].append(f"D{r}")
        r = ws.max_row + 1
        ws.append(["", f"Sub-total – {kind}", "", "", "", f"=SUBTOTAL(9,F{top}:F{r - 1})"])
        ws.cell(r, 6).number_format = F_INR
        bold_row(ws, r, SUB_FILL)
    last = ws.max_row
    r = last + 1
    ws.append(["", "GRAND TOTAL", "", "", "", f"=SUBTOTAL(9,F{first}:F{last})"])
    ws.cell(r, 6).number_format = F_INR
    bold_row(ws, r, INPUT_FILL)
    total = r
    ws.append(["", "Grand total in lakh (₹ lakh)", "", "", "", f"=F{total}/100000"])
    ws.cell(ws.max_row, 6).number_format = "#,##0.00"
    ws.append([])
    ws.append(["", "Abstract of quantities"])
    bold_row(ws, ws.max_row)
    summary = [("Total RCC concrete", "m³", "rcc"), ("Total PCC", "m³", "pcc")]
    summary += [("Total reinforcement steel", "kg", "steel"), ("Total centering & shuttering", "m²", "formwork")]
    at: dict[str, int] = {}
    for label, unit, head in summary:
        if not cells[head]:
            continue
        ws.append(["", label, unit, f"=SUM({','.join(cells[head])})"])
        ws.cell(ws.max_row, 4).number_format = UNIT_FMT[unit]
        at[head] = ws.max_row
    if "rcc" in at and "steel" in at:
        c, s = f"D{at['rcc']}", f"D{at['steel']}"
        ws.append(["", "Steel / RCC concrete ratio", "kg/m³", f"=IF({c}=0,0,{s}/{c})"])
        ws.cell(ws.max_row, 4).number_format = "0.0"
        ws.append(["", "Cost per m³ of RCC (all-in)", "₹/m³", f"=IF({c}=0,0,F{total}/{c})"])
        ws.cell(ws.max_row, 4).number_format = F_INR
    footnote(
        ws,
        "Rates are taken from the Rates sheet (yellow input cells); amounts and totals are live formulas.",
        "Beam concrete excludes the slab depth and the beam–column joint (counted with the slab / column). "
        "Column and wall steel includes 10 % for laps.",
    )
    page_setup(ws, hrow, [9, 62, 7, 13, 13, 16])
    from openpyxl.styles import Alignment

    for row in ws.iter_rows(min_row=first, max_row=last, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")


def _breakdown_sheet(wb, title: str, head: str, groups: dict, grades: list[str], refs: dict, project, wm) -> None:
    """Rows = floors / member types; per-grade concrete, steel, formwork and formula columns."""
    from openpyxl.utils import get_column_letter as L

    ws = wb.create_sheet(title)
    hrow = header_block(ws, f"Quantities and cost {title.lower()}", project, wm)
    rcc = [g for g in grades if not g.startswith("PCC")]
    headers = [head, *(f"{'RCC ' if g in rcc else ''}{g} m³" for g in grades)]
    headers += ["Steel kg", "Formwork m²", "RCC total m³", "Steel kg/m³", "Cost ₹"]
    table_header(ws, hrow, headers)
    ng = len(grades)
    cs, cf, ct, _, cc = (L(ng + k) for k in range(2, 7))
    rcc_cols = [L(2 + grades.index(g)) for g in rcc]
    first = hrow + 1
    for name, v in groups.items():
        if not (v.get("by_grade") or v.get("steel") or v.get("formwork")):
            continue
        r = ws.max_row + 1
        price = [f"{L(2 + i)}{r}*{refs[rate_key(g)]}" for i, g in enumerate(grades)]
        price += [f"{cs}{r}*{refs['steel_kg']}", f"{cf}{r}*{refs['formwork_m2']}"]
        ws.append(
            [
                name,
                *(round(v.get("by_grade", {}).get(g, 0.0), 3) for g in grades),
                round(v.get("steel", 0.0), 3),
                round(v.get("formwork", 0.0), 3),
                f"=SUM({','.join(f'{c}{r}' for c in rcc_cols)})" if rcc_cols else 0.0,
                f"=IF({ct}{r}=0,0,{cs}{r}/{ct}{r})",
                "=" + "+".join(price),
            ]
        )
    last = ws.max_row
    r = last + 1
    tot = ["TOTAL"] + [f"=SUM({L(i)}{first}:{L(i)}{last})" for i in range(2, ng + 5)]
    ws.append([*tot, f"=IF({ct}{r}=0,0,{cs}{r}/{ct}{r})", f"=SUM({cc}{first}:{cc}{last})"])
    bold_row(ws, r, SUB_FILL)
    for row in ws.iter_rows(min_row=first, max_row=r):
        for i, c in enumerate(row[1:], 2):
            c.number_format = F_KG if i == ng + 2 else "0.0" if i == ng + 5 else F_INR if i == ng + 6 else F_M3
    footnote(ws, "Cost = quantities × rates on the Rates sheet (formulas). PCC is not part of the RCC total.")
    page_setup(ws, hrow, [16] + [13] * (len(headers) - 1), landscape=len(headers) > 8)


def _revisions_sheet(wb, project: Project, revs: list[dict], watermark: str) -> None:
    ws = wb.create_sheet("Revisions")
    hrow = header_block(ws, "BOQ revisions", project, watermark)
    table_header(ws, hrow, ["Revision", "Date", "Concrete m³", "Steel kg", "Formwork m²", "Cost ₹"])
    for rv in revs:
        t = rv.get("total", {})
        ws.append([rv.get("label", ""), rv.get("date", "")] + [t.get(q, 0.0) for q in ("concrete", "steel")])
        r = ws.max_row
        ws.cell(r, 5, t.get("formwork", 0.0))
        ws.cell(r, 6, t.get("cost", 0.0))
        for col, fmt in zip(range(3, 7), (F_M3, F_KG, F_M2, F_INR)):
            ws.cell(r, col).number_format = fmt
    if len(revs) >= 2:
        a, b = revs[-2], revs[-1]
        ws.append([])
        ws.append([f"Comparison {a.get('label', 'A')} → {b.get('label', 'B')}"])
        bold_row(ws, ws.max_row)
        r0 = ws.max_row + 1
        heads = ["Group", "Item", f"A: {a.get('label', '')}", f"B: {b.get('label', '')}", "Change", "Change %"]
        table_header(ws, r0, heads)
        for g, item, va, vb, _ch, _pct in compare_revisions(a, b):
            r = ws.max_row + 1
            ws.append([g, item, va, vb, f"=D{r}-C{r}", f'=IF(C{r}=0,"",E{r}/C{r})'])
            fmt = F_INR if "(₹)" in item else F_KG if "(kg)" in item else F_M3
            for col in (3, 4, 5):
                ws.cell(r, col).number_format = fmt
            ws.cell(r, 6).number_format = F_PCT
    footnote(ws, "Revisions are saved from the BOQ in the app ('save revision R1').")
    page_setup(ws, hrow, [16, 34, 16, 16, 14, 12])


def write_boq_excel(path: str, project: Project, rep: DesignReport, watermark: str = "") -> str:
    """Write the editable BOQ / cost estimate workbook (see module docs) and return ``path``."""
    from openpyxl import Workbook

    boq = rep.boq or {}
    if not boq.get("by_type"):
        raise ValueError("no quantities – run the design first")
    wb = Workbook()
    wb.remove(wb.active)
    refs = _rates_sheet(wb, project, boq, watermark)
    _boq_sheet(wb, project, boq, refs, watermark)
    grades = _ordered_grades(boq.get("concrete", {}))
    _breakdown_sheet(wb, "By floor", "Floor", boq.get("by_level", {}), grades, refs, project, watermark)
    _breakdown_sheet(wb, "By type", "Member type", boq["by_type"], grades, refs, project, watermark)
    revs = (project.meta or {}).get("revisions") or []
    if revs:
        _revisions_sheet(wb, project, revs, watermark)
    wb.active = 0
    wb.save(path)
    return path
