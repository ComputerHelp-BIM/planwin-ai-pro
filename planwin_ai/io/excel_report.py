"""Excel workbook of plan checks, analysis results, member designs and BOQ."""

from __future__ import annotations

import datetime as _dt

from .. import APP_NAME, COMPANY, __version__
from ..core.beamcalc import eighth_points
from ..core.frame import FrameAnalysis, FrameModel
from ..core.model import Project
from ..core.plan_engine import PlanResult
from ..design.report import DesignReport
from ..units import Units
from .report_common import DISCLAIMER, literal_text

#: kN based unit symbols in report headers, most specific first
_HEADER_UNITS = (
    ("kN·m", "moment"),
    ("kNm", "moment"),
    ("kN/m³", "unit_weight"),
    ("kN/m²", "area"),
    ("kN/m2", "area"),
    ("kN/m", "line"),
    ("kN", "force"),
)


def header_quantity(header: str) -> str | None:
    """The converted quantity (force, moment, ...) a report column header is expressed in, if any."""
    for sym, q in _HEADER_UNITS:
        if sym in str(header):
            return q
    return None


def convert_rows(headers, rows, units: Units):
    """Headers and rows in display units: kN based columns (detected from the header) are converted."""
    qty = [header_quantity(h) for h in headers]
    if not units.mks or not any(qty):
        return list(headers), [tuple(r) for r in rows]
    out = []
    for r in rows:
        out.append(
            tuple(
                units.show(float(v), q) if q and isinstance(v, (int, float)) and not isinstance(v, bool) else v
                for v, q in zip(r, qty)
            )
            + tuple(r[len(qty) :])
        )
    return [units.text(h) for h in headers], out


def seismic_method_summary(project: Project, fm: FrameModel | None, fa: FrameAnalysis | None, rep) -> dict:
    """Seismic analysis method used and why (IS 1893-1:2016 cl 7.7.1), with modal / RS results."""
    s = project.seismic
    rs = dict(getattr(fa, "rs", None) or getattr(fm, "rs", None) or {})
    modal = getattr(fa, "modal", None) or getattr(rep, "modal", None)
    irr = list(getattr(fa, "irregularities", None) or getattr(rep, "irregularities", None) or [])
    if fa is not None:
        required = bool(fa.dynamic_required)
    else:
        from ..core.dynamics import dynamic_analysis_required
        from ..core.irregularity import is_regular

        if fm is not None and fm.levels:
            base = fm.levels[s.base_level].z if s.base_level < len(fm.levels) else 0.0
            height = fm.levels[-1].z - base
        else:
            el = project.elevations()
            height = el[-1] - (el[s.base_level] if s.base_level < len(el) else 0.0)
        required = s.enabled and dynamic_analysis_required(project, height, is_regular(irr))
    if not s.enabled:
        method, why = "none", "Seismic loads are disabled for this project"
    elif rs:
        method = "Response spectrum (IS 1893-1:2016 cl 7.7.3, CQC)"
        if s.method == "response_spectrum":
            why = "Response spectrum method selected by the user"
        else:
            why = "cl 7.7.1: dynamic analysis required (not a regular building below 15 m in zone II)"
    else:
        method = "Equivalent static (IS 1893-1:2016 cl 7.6)"
        if s.method == "response_spectrum" or (s.method == "auto" and required):
            why = "Response spectrum analysis was not possible – equivalent static method used (see warnings)"
        elif required:
            why = "WARNING: cl 7.7.1 requires dynamic analysis – static method selected by the user"
        elif s.method == "static":
            why = "Equivalent static method selected by the user; permitted by cl 7.7.1"
        else:
            why = "cl 7.7.1: regular building below 15 m in zone II – equivalent static method permitted"
    return {
        "method": method,
        "why": why,
        "setting": s.method,
        "required": required,
        "rs": rs,
        "modal": modal,
        "irregular": any(bool(i.irregular) for i in irr),
    }


def storey_shear_rows(fm: FrameModel | None, rs: dict) -> list[tuple]:
    """(level, static VX, static VY, RSX, RSY) per level index, top first – scaled RS storey shears."""
    if fm is None:
        return []
    rows = []
    for i in range(len(fm.levels) - 1, -1, -1):
        row: list = [fm.levels[i].name]
        for d in ("X", "Y"):
            sr = fm.seismic.get("EQ" + d)
            row.append(float(sum(sr.forces[i:])) if sr else "")
        for c in ("RSX", "RSY"):
            r = rs.get(c)
            row.append(float(r.storey_shear[i]) if r and i < len(r.storey_shear) else "")
        rows.append(tuple(row))
    return rows


def boq_rows(groups: dict) -> list[tuple]:
    """(name, concrete m³, PCC m³, steel kg, formwork m², cost, steel kg/m³) per BOQ group + TOTAL."""
    rows = []
    tot = dict.fromkeys(("concrete", "pcc", "steel", "formwork", "cost"), 0.0)
    for k, v in groups.items():
        c = v.get("concrete", 0.0)
        rows.append(
            (
                k,
                c,
                v.get("pcc", 0.0),
                v.get("steel", 0.0),
                v.get("formwork", 0.0),
                v.get("cost", 0.0),
                v.get("steel", 0.0) / c if c else 0.0,
            )
        )
        for q in tot:
            tot[q] += v.get(q, 0.0)
    if rows:
        c = tot["concrete"]
        rows.append(
            ("TOTAL", c, tot["pcc"], tot["steel"], tot["formwork"], tot["cost"], tot["steel"] / c if c else 0.0)
        )
    return rows


BOQ_HEADERS = ["Concrete m³", "PCC m³", "Steel kg", "Formwork m²", "Cost ₹", "Steel kg/m³"]


def ductile_summary(rep) -> dict:
    checks = [(d, name, ok, det) for d in rep.ductile for name, ok, det in d.checks]
    return {
        "members": len(rep.ductile),
        "members_failing": sum(not d.ok for d in rep.ductile),
        "checks": len(checks),
        "passed": sum(1 for c in checks if c[2]),
        "failed": sum(1 for c in checks if not c[2]),
        "failures": [c for c in checks if not c[2]],
    }


def write_excel(
    path: str,
    project: Project,
    plan_results: dict[str, PlanResult],
    fm: FrameModel | None = None,
    fa: FrameAnalysis | None = None,
    rep: DesignReport | None = None,
    watermark: str = "",
    units: Units | None = None,
) -> str:
    """Write the workbook.  ``units`` (default: the app-wide :data:`planwin_ai.units.current`)
    sets the display units of force, moment, line-load and pressure columns."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if units is None:
        from .. import units as _units

        units = _units.current
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(color="FFFFFF", bold=True)
    bad_fill = PatternFill("solid", fgColor="F8D7DA")
    sub_font = Font(bold=True, color="1F3A5F")

    def _cell(v):
        if isinstance(v, float):
            return round(v, 3) if v == v and abs(v) != float("inf") else str(v)
        return v

    def style_header(ws, row: int):
        for c in ws[row]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def sheet(title, headers, rows, bad_col: int | None = None, pre: list | None = None):
        """Sheet with a header row; ``pre`` rows (e.g. a summary) are written above the header."""
        headers, rows = convert_rows(headers, rows, units)
        ws = wb.create_sheet(title[:31])
        for r in pre or []:
            ws.append([_cell(v) for v in r])
            if r and len(r) == 1:
                ws.cell(ws.max_row, 1).font = sub_font
        ws.append(headers)
        hrow = ws.max_row
        style_header(ws, hrow)
        for r in rows:
            ws.append([_cell(v) for v in r])
            if bad_col is not None and r[bad_col] in (False, "FAIL", "NO"):
                for c in ws[ws.max_row]:
                    c.fill = bad_fill
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r)] + [8])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 45)
        ws.freeze_panes = f"A{hrow + 1}"
        return ws

    def block(ws, title, headers, rows):
        """A titled table appended below the existing content of ``ws``."""
        headers, rows = convert_rows(headers, rows, units)
        if ws.max_row > 1 or ws.cell(1, 1).value is not None:
            ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = sub_font
        ws.append(headers)
        style_header(ws, ws.max_row)
        for r in rows:
            ws.append([_cell(v) for v in r])

    ws = wb.active
    ws.title = "Summary"
    ws.append([f"{APP_NAME} {__version__} – {COMPANY}"])
    ws["A1"].font = Font(size=14, bold=True)
    info = [
        ("Project", project.name),
        ("Client", project.client),
        ("Engineer", project.engineer),
        ("Location", project.location),
        ("Date", _dt.date.today().isoformat()),
        ("Seismic zone", project.seismic.zone),
        ("Wind city / Vb", f"{project.wind.city} / {project.wind.basic_speed} m/s"),
        ("Height (m)", round(project.elevations()[-1], 3)),
        ("Levels", len(project.levels)),
    ]
    sm = seismic_method_summary(project, fm, fa, rep)
    info.append(("Seismic method", sm["method"]))
    info.append(("Units", "MKS (t, m)" if units.mks else "SI (kN, m)"))
    if watermark:
        info.append(("Licence", watermark))
    for k, v in info:
        ws.append([k, v])
    ws.append([])
    ws.append([DISCLAIMER])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 50

    rows = []
    for pname, pr in plan_results.items():
        rows.append(
            (
                pname,
                pr.applied["D"],
                pr.applied["L"],
                pr.reacted["D"],
                pr.reacted["L"],
                pr.imbalance_pct,
                len(pr.errors),
                len(pr.warnings),
            )
        )
    sheet(
        "Plan checks",
        [
            "Plan",
            "Applied DL kN",
            "Applied LL kN",
            "Reaction DL kN",
            "Reaction LL kN",
            "Imbalance %",
            "Errors",
            "Warnings",
        ],
        rows,
    )
    rows = []
    for pname, pr in plan_results.items():
        for cl in pr.columns.values():
            rows.append((pname, cl.mark, cl.dead, cl.live, cl.total))
    sheet("Column loads (plan)", ["Plan", "Column", "Dead kN", "Live kN", "Total kN"], rows)
    rows = []
    for pname, pr in plan_results.items():
        plan = project.plan(pname)
        bmap = {b.id: b for b in plan.beams} if plan else {}
        for bid, br in pr.beams.items():
            b = bmap.get(bid)
            if not b:
                continue
            dia = br.diagram(True)
            pts = eighth_points(br.length, dia)
            rows.append(
                (
                    pname,
                    b.mark,
                    br.length,
                    b.b,
                    b.d,
                    br.total("D"),
                    br.total("L"),
                    br.equivalent_udl(),
                    *[p["M"] for p in pts],
                    max(abs(p["V"]) for p in pts),
                )
            )
    sheet(
        "Beam loads (plan)",
        [
            "Plan",
            "Beam",
            "Span m",
            "b m",
            "D m",
            "Total DL kN",
            "Total LL kN",
            "Eq. UDL kN/m",
            *[f"Mu @{k}/8 kNm" for k in range(9)],
            "Vu max kN",
        ],
        rows,
    )
    issues = []
    for pname, pr in plan_results.items():
        for i in pr.issues:
            issues.append((pname, i.level, i.kind, i.message))
    if fm:
        for i in fm.issues:
            issues.append(("Frame", i.level, i.kind, i.message))
    sheet("Issues", ["Where", "Level", "Type", "Message"], issues)
    if fa and fm:
        rows = []
        for nid, nd in fm.nodes.items():
            if not nd.support:
                continue
            for c in fa.combos:
                r = fa.reaction(nid, c.factors)
                rows.append((nd.tag, c.name, *[float(v) for v in r]))
        sheet(
            "Support reactions",
            ["Column", "Combination", "FX kN", "FY kN", "FZ kN", "MX kNm", "MY kNm", "MZ kNm"],
            rows,
        )
        rows = [(k, r.T, r.sa_g, r.Ah, r.W, r.Vb) for k, r in fm.seismic.items()]
        if rows:
            sheet("Seismic", ["Direction", "T s", "Sa/g", "Ah", "W kN", "VB kN"], rows)
        rows = []
        for k, r in fm.wind.items():
            for i, f in enumerate(r.forces):
                if f:
                    rows.append((k, fm.levels[i].name, r.pressures[i], f))
        if rows:
            sheet("Wind", ["Direction", "Level", "pd kN/m2", "Storey force kN"], rows)
    if project.seismic.enabled and (fm or rep):
        modal = sm["modal"]
        pre = [
            ("Seismic analysis method (IS 1893-1:2016)",),
            ("Method used", sm["method"]),
            ("Why", sm["why"]),
            ("Method setting", sm["setting"]),
            ("Dynamic analysis required (cl 7.7.1)", "YES" if sm["required"] else "NO"),
            ("Irregular configuration (Tables 5/6)", "YES" if sm["irregular"] else "NO"),
            (),
            ("Modal periods and effective modal mass (cl 7.7.5.2)",),
        ]
        rows = []
        if modal is not None:
            cx, cy = modal.cumulative("x"), modal.cumulative("y")
            for i, m in enumerate(modal.modes):
                rows.append(
                    (
                        m.number,
                        m.period,
                        100 * m.mass_x,
                        100 * m.mass_y,
                        100 * m.mass_rz,
                        100 * cx[i],
                        100 * cy[i],
                    )
                )
        ws = sheet(
            "Seismic method",
            ["Mode", "T s", "Mass X %", "Mass Y %", "Mass RZ %", "Cumulative X %", "Cumulative Y %"],
            rows,
            pre=pre,
        )
        ws.column_dimensions["A"].width = 38
        if sm["rs"]:
            block(
                ws,
                "Response spectrum scaling to the static base shear (cl 7.7.3)",
                ["Direction", "Modes used", "VB dynamic kN", "VB static kN", "Scale factor"],
                [(k, r.modes_used, r.vb_dynamic, r.vb_static, r.scale) for k, r in sm["rs"].items()],
            )
        block(
            ws,
            "Storey shears (response spectrum scaled)" if sm["rs"] else "Storey shears (equivalent static)",
            ["Level", "Static VX kN", "Static VY kN", "RSX kN", "RSY kN"],
            storey_shear_rows(fm, sm["rs"]),
        )
    if rep and rep.irregularities:
        sheet(
            "Irregularity",
            ["Table", "Irregularity", "Clause", "Regular", "Status", "Detail"],
            [
                (
                    i.table,
                    i.name,
                    i.clause,
                    "-" if i.irregular is None else ("NO" if i.irregular else "YES"),
                    i.status,
                    i.detail,
                )
                for i in rep.irregularities
            ],
            bad_col=3,
        )
    if rep:
        sheet(
            "Beam design",
            [
                "Level",
                "Beam",
                "b m",
                "D m",
                "Span m",
                "Mu+ kNm",
                "Mu- left kNm",
                "Mu- right kNm",
                "Vu kN",
                "Bottom",
                "Top left",
                "Top right",
                "Stirrups",
                "End-zone links (2d)",
                "Tu kNm",
                "Side face",
                "OK",
                "Notes",
            ],
            [
                (
                    b.level,
                    b.mark,
                    b.b,
                    b.d,
                    b.span,
                    b.M_sag,
                    b.M_hog_l,
                    b.M_hog_r,
                    b.V_max,
                    b.bottom,
                    b.top_l,
                    b.top_r,
                    b.stirrups,
                    str(b.links_end) if b.links_end else "",
                    b.T_max,
                    b.side_face,
                    "YES" if b.ok else "NO",
                    "; ".join(b.notes),
                )
                for b in rep.beams
            ],
            bad_col=16,
        )
        sheet(
            "Column design",
            [
                "Level",
                "Column",
                "b m",
                "D m",
                "Pu kN",
                "Mux kNm",
                "Muy kNm",
                "Steel %",
                "Bars",
                "Ties",
                "Confining hoops (l0)",
                "l0 m",
                "Governing",
                "Interaction",
                "OK",
                "Notes",
            ],
            [
                (
                    c.level,
                    c.mark,
                    c.b,
                    c.d,
                    c.Pu,
                    c.Mux,
                    c.Muy,
                    c.steel_pct,
                    c.bars,
                    c.ties,
                    str(c.tie_confined) if c.tie_confined else "",
                    c.l0 if c.tie_confined else "",
                    c.governing,
                    c.utilisation,
                    "YES" if c.ok else "NO",
                    "; ".join(c.notes),
                )
                for c in rep.columns
            ],
            bad_col=14,
        )
        sheet(
            "Footing design",
            ["Column", "P service kN", "L m", "B m", "D m", "Bars along L", "Bars along B", "q kN/m2", "OK", "Notes"],
            [
                (
                    f.mark,
                    f.P_service,
                    f.L,
                    f.B,
                    f.D,
                    f.bars_L,
                    f.bars_B,
                    f.q,
                    "YES" if f.ok else "NO",
                    "; ".join(f.notes),
                )
                for f in rep.footings
            ],
            bad_col=8,
        )
        sheet(
            "Slab design",
            [
                "Plan",
                "Slab",
                "lx m",
                "ly m",
                "Type",
                "D mm",
                "Mx+ kNm/m",
                "My+ kNm/m",
                "M- kNm/m",
                "Short span",
                "Long span",
                "Support (top)",
                "OK",
                "Notes",
            ],
            [
                (
                    pn,
                    s.mark,
                    s.lx,
                    s.ly,
                    s.kind,
                    s.D_mm,
                    s.Mx_pos,
                    s.My_pos,
                    s.M_neg,
                    s.ast_x,
                    s.ast_y,
                    s.ast_neg,
                    "YES" if s.ok else "NO",
                    "; ".join(s.notes),
                )
                for pn, s in rep.slabs
            ],
            bad_col=12,
        )
        sheet(
            "Drift",
            ["Case", "Level", "Drift mm", "Ratio", "OK (<=0.004)"],
            [(d["case"], d["level"], d["drift_mm"], d["ratio"], "YES" if d["ok"] else "NO") for d in rep.drifts],
            bad_col=4,
        )
        boq = rep.boq
        sheet(
            "BOQ & Cost",
            ["Item", "Unit", "Quantity", "Rate (INR)", "Amount (INR)"],
            list(boq.get("lines", []))
            + [
                ("TOTAL", "", "", "", boq.get("cost", 0.0)),
                ("Steel / concrete", "kg/m³", boq.get("steel_per_m3", 0.0), "", ""),
            ],
        )
        if boq.get("by_level"):
            sheet("BOQ by floor", ["Floor", *BOQ_HEADERS], boq_rows(boq["by_level"]))
        if boq.get("by_type"):
            sheet("BOQ by type", ["Member type", *BOQ_HEADERS], boq_rows(boq["by_type"]))
        if rep.walls:
            sheet(
                "Walls",
                [
                    "Level",
                    "Wall",
                    "t m",
                    "Lw m",
                    "Height m",
                    "Pu kN",
                    "Mu kNm",
                    "Vu kN",
                    "ρv",
                    "ρh",
                    "Curtains",
                    "Vertical",
                    "Horizontal",
                    "Boundary elements",
                    "Utilisation",
                    "OK",
                    "Notes",
                ],
                [
                    (
                        w.level,
                        w.mark,
                        w.t,
                        w.Lw,
                        w.height,
                        w.Pu,
                        w.Mu,
                        w.Vu,
                        round(w.rho_v, 5),
                        round(w.rho_h, 5),
                        w.curtains,
                        w.vertical,
                        w.horizontal,
                        w.boundary,
                        w.utilisation,
                        "YES" if w.ok else "NO",
                        "; ".join(w.notes),
                    )
                    for w in rep.walls
                ],
                bad_col=15,
            )
        if rep.combined_footings:
            sheet(
                "Combined footings",
                [
                    "Columns",
                    "P service kN",
                    "L m",
                    "B m",
                    "D m",
                    "x m",
                    "y m",
                    "Angle °",
                    "q kN/m²",
                    "Top (between columns)",
                    "Bottom",
                    "Transverse",
                    "OK",
                    "Notes",
                ],
                [
                    (
                        " + ".join(c.marks),
                        c.P_service,
                        c.L,
                        c.B,
                        c.D,
                        c.x,
                        c.y,
                        c.angle,
                        c.q,
                        c.top,
                        c.bottom,
                        "; ".join(c.transverse),
                        "YES" if c.ok else "NO",
                        "; ".join(c.notes),
                    )
                    for c in rep.combined_footings
                ],
                bad_col=12,
            )
        if rep.ductile:
            ds = ductile_summary(rep)
            pre = [
                ("IS 13920:2016 ductile detailing checks",),
                ("Members checked", ds["members"]),
                ("Members failing", ds["members_failing"]),
                ("Checks", ds["checks"]),
                ("Passed", ds["passed"]),
                ("Failed", ds["failed"]),
                (),
            ]
            sheet(
                "IS 13920",
                ["Kind", "Member", "Level", "Check", "Result", "Detail"],
                [
                    (d.kind, d.mark, d.level, name, "PASS" if ok else "FAIL", det)
                    for d in rep.ductile
                    for name, ok, det in d.checks
                ],
                bad_col=4,
                pre=pre,
            )
            sheet(
                "IS 13920 detailing",
                ["Kind", "Member", "Level", "Detailing"],
                [(d.kind, d.mark, d.level, line) for d in rep.ductile for line in d.detailing],
            )
    revs = (project.meta or {}).get("revisions") or []
    if len(revs) >= 2:
        from ..design.quantities import compare_revisions

        a, b = revs[-2], revs[-1]
        sheet(
            "Revision compare",
            [
                "Group",
                "Item",
                f"A: {a.get('label', '')} {a.get('date', '')}".strip(),
                f"B: {b.get('label', '')} {b.get('date', '')}".strip(),
                "Change",
                "Change %",
            ],
            [
                (g, item, va, vb, ch, pct if abs(pct) != float("inf") else "new")
                for g, item, va, vb, ch, pct in compare_revisions(a, b)
            ],
        )
    literal_text(wb)  # names/marks starting with '=' stay text
    wb.save(path)
    return path
