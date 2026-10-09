"""Bar bending schedule (BBS) from the structured detailing of a design run.

The schedule is built from :class:`~planwin_ai.design.report.DesignReport` (bar sets, links and
meshes chosen by the IS 456 design routines), the analysed :class:`~planwin_ai.core.frame.FrameModel`
(member connectivity, support widths) and the project's design settings (covers, fy).
All internal dimensions are in mm; cutting lengths are reported in m.

Detailing rules (IS 456:2000, SP 34:1987)
-----------------------------------------
* **Development length** ``Ld = φ·σs / (4·τbd)`` with ``σs = 0.87 fy`` and ``τbd`` from IS 456
  cl 26.2.1.1 (M20 1.2, M25 1.4, M30 1.5, M35 1.7, M40+ 1.9 N/mm²; 1.0 below M20) increased by
  60 % for deformed bars (cl 26.2.1.1, note).  In compression ``τbd`` is further increased by 25 %.
* **Lap length** in flexural tension ``max(Ld, 30φ)`` (cl 26.2.5.1 c); for the usual grades this is
  simply ``Ld`` (≈ 48φ for M25 / Fe 500).  Compression laps (hanger bars) ``max(Ld,c, 24φ)``
  (cl 26.2.5.1 d).  Column bars use the tension lap (conservative under lateral load).
* **Bend deductions** (SP 34): 45° = 1d, 90° = 2d, 135° = 3d, 180° = 4d, applied to out-to-out
  dimensions.  Anchorage value of a standard 90° bend = 8d (incl. its 4d extension), U hook = 16d.
* **Anchorage leg**: where the straight length available inside a support is less than ``Ld`` the
  bar is bent through 90°; with the 8d bend value the leg is ``max(Ld − straight − 4d, 12d)``.
* **Links / ties**: closed rectangular links of outer size ``(b − 2c) × (D − 2c)`` with two 135°
  hooks of ``max(10d, 75 mm)`` extension:
  ``L = 2(b − 2c) + 2(D − 2c) + 2·hook − (3 × 2d + 2 × 3d)``.  Number ``= ceil(L/s) + 1``.
  Four-legged links are scheduled as the outer link plus an inner closed link one third as wide.
* **Beams** (one plan beam = one physical beam; frame segments of the same ``group`` at the same
  level are joined and their spans summed):

  - bottom bars (largest bar set of the segments) run the full length between support centre-lines
    and continue to the far face of the end supports less cover, then bend up (90°) for ``Ld``
    measured from the inner face; at a free (cantilever) end they stop at the cover;
  - top support bars extend ``0.3 × span`` into each adjacent span (support centre-line reference);
    at an end support they continue to the far face and bend down for ``Ld``; over a cantilever
    they run to the tip and bend down by ``D − 2c``;
  - 2 hanger bars (T12, or the top bar diameter if smaller) span the middle zone between the
    top-bar cut-offs with a compression lap at each end;
  - legs at supports formed by another beam are limited to ``D − 2c``.
* **Columns** (per storey): main bars = storey height + lap (lap at each floor, placed on the bar
  of the storey below).  Bottom storey bars continue ``D − cover`` into the footing and end in a
  90° foot ``max(Ld − (Df − cf) − 4d, 12d)``; top storey bars stop at the roof less cover and are
  bent into the roof beam for ``Ld`` measured from its soffit (min 12d).  Ties as links above.
* **Footings**: bars along L and along B with end cover; number ``= ceil((width − 2c)/s) + 1``;
  both ends bent up by ``D − 2c`` (U-bar).  Identical footings are grouped.
* **Slabs** (rectangular panels, bars between support centre-lines): bottom bars both ways at the
  designed spacing (``mesh_x`` short span, ``mesh_y`` long span / distribution); where ``mesh_neg``
  exists, top bars over each continuous edge shared by two panels extend ``0.25 lx`` into each
  panel, with distribution bars (``mesh_y``) under them.  Cantilever slabs get top main bars over
  the cantilever span plus ``Ld`` into the back span (bent down at the tip) and top distribution.
* Bars longer than the 12 m stock length are scheduled with additional tension laps.

Not included: wastage, chairs/spacers, IS 13920 confining links, slab corner torsion steel
(Annex D-1.8), nominal top steel at discontinuous slab edges and anchorage of floating columns.
Shear walls and combined footings are not scheduled; each one is listed in the warnings
(column bars still get their foot in a combined footing).

Shape codes
-----------
Numeric shape codes in the IS 2502 / BS 8666 convention (dimensions out-to-out, mm):

* ``00`` straight bar – ``a`` = length;
* ``11`` one 90° bend (L-bar) – ``a`` = main length, ``b`` = leg;
* ``21`` two 90° bends to the same side (U-bar) – ``a`` and ``c`` = legs, ``b`` = main length;
* ``51`` closed rectangular link with 135° hooks – ``a`` = width, ``b`` = height, ``c`` = hook.
"""

from __future__ import annotations

import datetime as _dt
import math
import re
from collections import defaultdict
from dataclasses import dataclass, field

from .. import APP_NAME, COMPANY, __version__
from ..core import geometry as G
from ..core.frame import FrameModel
from ..core.model import Project, grade_fck
from ..design.is456.common import BarMesh, BarSet, Links
from ..design.report import BeamDesign, DesignReport
from .report_common import DISCLAIMER, literal_text

STOCK_LENGTH = 12000.0  # mm
CUT_NOTE = "Verify against approved structural drawings before cutting"
SHAPE_CODES = {
    "00": "Straight bar (a)",
    "11": "One 90° bend, L-bar (a, b = leg)",
    "21": "Two 90° bends, U-bar (a, c = legs, b = length)",
    "51": "Closed rectangular link with 135° hooks (a × b, c = hook)",
}
BEND_DEDUCTION = {45: 1.0, 90: 2.0, 135: 3.0, 180: 4.0}  # × bar diameter (SP 34)
BEND_ANCHORAGE_90 = 8.0  # anchorage value of a standard 90° bend, × d (SP 34)
U_HOOK_ANCHORAGE = 16.0  # anchorage value of a standard U hook, × d (SP 34)
MEMBER_TYPES = ("Beam", "Column", "Footing", "Slab")


# ---------------------------------------------------------------- data model
@dataclass
class BarItem:
    member_type: str  # "Beam" | "Column" | "Footing" | "Slab"
    member: str
    level: str
    bar_mark: str
    description: str
    shape_code: str
    dia: int  # mm
    count: int  # bars of this mark in ONE member
    members: int  # identical members
    cutting_length: float  # m
    dims: dict[str, float] = field(default_factory=dict)  # mm

    @property
    def total_count(self) -> int:
        return self.count * self.members

    @property
    def total_length(self) -> float:
        return self.cutting_length * self.total_count

    @property
    def unit_weight(self) -> float:
        return unit_weight(self.dia)

    @property
    def weight(self) -> float:
        return self.total_length * self.unit_weight


@dataclass
class BBS:
    items: list[BarItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    concrete_m3: float = 0.0  # structural concrete (from the design BOQ), for kg/m³
    boq_steel: float = 0.0  # steel estimate of the design BOQ (kg), for comparison

    def _sum_by(self, key) -> dict:
        out: dict = defaultdict(float)
        for it in self.items:
            out[key(it)] += it.weight
        return dict(out)

    def by_dia(self) -> dict[int, float]:
        return dict(sorted(self._sum_by(lambda it: it.dia).items()))

    def by_level(self) -> dict[str, float]:
        return self._sum_by(lambda it: it.level)

    def by_member_type(self) -> dict[str, float]:
        return self._sum_by(lambda it: it.member_type)

    @property
    def total_weight(self) -> float:
        return sum(it.weight for it in self.items)

    @property
    def steel_per_m3(self) -> float:
        return self.total_weight / self.concrete_m3 if self.concrete_m3 > 0 else 0.0


# ---------------------------------------------------------------- detailing rules
def unit_weight(dia: float) -> float:
    """kg/m of a bar of diameter ``dia`` mm (d²/162)."""
    return dia * dia / 162.0


def bond_stress(fck: float) -> float:
    """Design bond stress τbd for plain bars in tension, IS 456 cl 26.2.1.1 (N/mm²)."""
    for grade, tbd in ((40, 1.9), (35, 1.7), (30, 1.5), (25, 1.4), (20, 1.2)):
        if fck >= grade:
            return tbd
    return 1.0


def development_length(dia: float, fck: float, fy: float, compression: bool = False) -> float:
    """Ld = φ·0.87fy / (4·τbd) for deformed bars (τbd +60 %, +25 % more in compression), mm."""
    tbd = bond_stress(fck) * 1.6 * (1.25 if compression else 1.0)
    return dia * 0.87 * fy / (4 * tbd)


def lap_length(dia: float, fck: float, fy: float, compression: bool = False) -> float:
    """Lap length (mm): max(Ld, 30φ) in flexural tension, max(Ld,c, 24φ) in compression."""
    if compression:
        return max(development_length(dia, fck, fy, True), 24 * dia)
    return max(development_length(dia, fck, fy), 30 * dia)


def hook_extension(dia: float) -> float:
    """Extension of a 135° link hook: 10d, at least 75 mm."""
    return max(10 * dia, 75.0)


def closed_link_length(a: float, b: float, dia: float) -> float:
    """Cutting length (mm) of a closed rectangular link of outer size a × b with two 135° hooks."""
    deduction = (3 * BEND_DEDUCTION[90] + 2 * BEND_DEDUCTION[135]) * dia
    return 2 * a + 2 * b + 2 * hook_extension(dia) - deduction


def anchorage_leg(Ld: float, straight: float, dia: float) -> float:
    """90° leg (mm) needed when only ``straight`` mm is available for anchorage ``Ld``; 0 if none."""
    rem = Ld - straight
    if rem <= 0:
        return 0.0
    return max(rem - (BEND_ANCHORAGE_90 - 4) * dia, 12 * dia)


def laps_needed(length: float, lap: float) -> int:
    """Number of laps for a bar of ``length`` mm supplied in 12 m stock lengths."""
    n = 0
    while length + n * lap > STOCK_LENGTH * (n + 1) and n < 100 and lap < STOCK_LENGTH:
        n += 1
    return n


def bent_bar(main: float, legs: list[float], dia: float) -> tuple[str, dict[str, float], float]:
    """Shape code, dims and cutting length (mm) of a bar with up to two 90° end legs."""
    legs = [lg for lg in legs if lg > 0]
    cut = main + sum(legs) - len(legs) * BEND_DEDUCTION[90] * dia
    if not legs:
        return "00", {"a": main}, cut
    if len(legs) == 1:
        return "11", {"a": main, "b": legs[0]}, cut
    return "21", {"a": legs[0], "b": main, "c": legs[1]}, cut


def _natural(s: str):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", s)]


def _ceil_count(length: float, spacing: float) -> int:
    return int(math.ceil(max(length, 0.0) / spacing - 1e-9)) + 1 if spacing > 0 else 0


def _larger(*sets):
    sets = [s for s in sets if s is not None and getattr(s, "count", 1) > 0]
    if not sets:
        return None
    return max(sets, key=lambda s: s.area if isinstance(s, BarSet) else s.area_per_m)


# ---------------------------------------------------------------- builder
class _Builder:
    def __init__(self, project: Project, fm: FrameModel, rep: DesignReport):
        self.p, self.fm, self.rep = project, fm, rep
        self.ds = project.design
        self.fy = self.ds.fy_main
        self.bbs = BBS(
            concrete_m3=float(rep.boq.get("total_concrete", 0.0) or 0.0),
            boq_steel=float(rep.boq.get("total_steel", 0.0) or 0.0),
        )
        self._marks: set[str] = set()
        self._seq: dict[str, int] = defaultdict(int)
        self.col_below: dict[int, object] = {}
        self.col_above: dict[int, object] = {}
        self.beams_at: dict[int, list] = defaultdict(list)
        for m in fm.members.values():
            if m.kind == "column":
                self.col_below[m.n2] = m
                self.col_above[m.n1] = m
            else:
                self.beams_at[m.n1].append(m)
                self.beams_at[m.n2].append(m)

    # ------------------------------------------------------------ helpers
    def _mark(self, prefix: str) -> str:
        while True:
            self._seq[prefix] += 1
            mk = f"{prefix}-{self._seq[prefix]:02d}"
            if mk not in self._marks:
                self._marks.add(mk)
                return mk

    def _add(self, mtype, member, level, prefix, desc, dia, count, cut_mm, shape, dims, members=1, lap=0.0):
        if count <= 0 or cut_mm <= 0 or dia <= 0:
            return
        n_laps = laps_needed(cut_mm, lap) if lap > 0 else 0
        if n_laps:
            cut_mm += n_laps * lap
            desc += f" (incl. {n_laps} lap{'s' if n_laps > 1 else ''} of {lap:.0f} mm)"
        self.bbs.items.append(
            BarItem(
                mtype,
                member,
                level,
                self._mark(prefix),
                desc,
                shape,
                int(dia),
                int(count),
                int(members),
                round(cut_mm) / 1000.0,
                {k: float(round(v)) for k, v in dims.items()},
            )
        )

    def _add_links(self, mtype, member, level, prefix, desc, links: Links, count, b_mm, d_mm, cover):
        a, h = b_mm - 2 * cover, d_mm - 2 * cover
        dims = {"a": a, "b": h, "c": hook_extension(links.dia)}
        self._add(mtype, member, level, prefix, desc, links.dia, count, closed_link_length(a, h, links.dia), "51", dims)
        inner = math.ceil(max(links.legs - 2, 0) / 2)
        if inner:
            ai = max(a / 3, 4 * links.dia)
            dims = {"a": ai, "b": h, "c": hook_extension(links.dia)}
            cut = closed_link_length(ai, h, links.dia)
            self._add(mtype, member, level, prefix, f"Inner {desc.lower()}", links.dia, count * inner, cut, "51", dims)

    def _support(self, node: int, member, group: str) -> tuple[str, float]:
        """('column' | 'beam' | 'free', half width of the support along the member in mm)."""
        a, b = self.fm.nodes[member.n1], self.fm.nodes[member.n2]
        L = math.hypot(b.x - a.x, b.y - a.y) or 1.0
        ux, uy = (b.x - a.x) / L, (b.y - a.y) / L
        col = self.col_below.get(node) or self.col_above.get(node)
        if col is not None:
            # same projection as design.quantities (half the column extent along the beam)
            ang = math.radians(col.angle)
            hw = (
                abs(ux * math.cos(ang) + uy * math.sin(ang)) * col.b / 2
                + abs(-ux * math.sin(ang) + uy * math.cos(ang)) * col.d / 2
            )
            return "column", hw * 1000
        others = [m for m in self.beams_at.get(node, []) if m.group != group]
        if others:
            return "beam", max(m.b for m in others) * 1000 / 2
        return "free", 0.0

    # ------------------------------------------------------------ beams
    def beams(self):
        groups: dict[tuple, list[BeamDesign]] = defaultdict(list)
        for bd in self.rep.beams:
            if bd.member_id in self.fm.members and bd.span > 0:
                groups[(bd.level_index, bd.group or f"#{bd.member_id}")].append(bd)
        for key in sorted(groups, key=lambda k: (k[0], _natural(groups[k][0].mark), k[1])):
            self._beam(groups[key])

    def _chain(self, segs: list[BeamDesign]) -> tuple[list[BeamDesign], list[int]]:
        """Order the segments of one plan beam end to end; returns (segments, node sequence)."""
        mems = {s.member_id: self.fm.members[s.member_id] for s in segs}
        if len(segs) == 1:
            m = mems[segs[0].member_id]
            return segs, [m.n1, m.n2]
        deg: dict[int, int] = defaultdict(int)
        for m in mems.values():
            deg[m.n1] += 1
            deg[m.n2] += 1
        ends = [n for n, k in deg.items() if k == 1]
        if len(ends) != 2:
            ordered = sorted(segs, key=lambda s: s.member_id)
            nodes = [mems[ordered[0].member_id].n1] + [mems[s.member_id].n2 for s in ordered]
            return ordered, nodes
        node, left, ordered, nodes = ends[0], list(segs), [], [ends[0]]
        while left:
            s = next((s for s in left if node in (mems[s.member_id].n1, mems[s.member_id].n2)), None)
            if s is None:
                break
            left.remove(s)
            m = mems[s.member_id]
            node = m.n2 if m.n1 == node else m.n1
            ordered.append(s)
            nodes.append(node)
        if left:  # disconnected – fall back to frame order
            ordered = sorted(segs, key=lambda s: s.member_id)
            nodes = [mems[ordered[0].member_id].n1] + [mems[s.member_id].n2 for s in ordered]
        return ordered, nodes

    def _beam(self, segs: list[BeamDesign]):
        segs, nodes = self._chain(segs)
        first = segs[0]
        mem0 = self.fm.members[first.member_id]
        mark, level, prefix = first.mark, first.level, f"{first.mark}-L{first.level_index}"
        fck = grade_fck(mem0.grade)
        c = self.ds.beam_cover * 1000
        D = first.d * 1000
        spans = [s.span * 1000 for s in segs]
        total = sum(spans)
        n = len(segs)
        sup = []
        for j, node in enumerate(nodes):
            seg = segs[min(j, n - 1)]
            sup.append(self._support(node, self.fm.members[seg.member_id], mem0.group))

        def end_detail(j: int, dia: float) -> tuple[float, float]:
            """(extension beyond the support centre-line, anchorage leg) at end node j."""
            kind, hw = sup[j]
            if kind == "free":
                return -c, 0.0
            leg = anchorage_leg(development_length(dia, fck, self.fy), 2 * hw - c, dia)
            if kind == "beam":
                leg = min(leg, D - 2 * c)
            return hw - c, leg

        # bottom bars – continuous over the whole plan beam
        bot = _larger(*(s.bottom_bars for s in segs))
        if bot is not None:
            e0, l0 = end_detail(0, bot.dia)
            e1, l1 = end_detail(n, bot.dia)
            shape, dims, cut = bent_bar(total + e0 + e1, [l0, l1], bot.dia)
            lap = lap_length(bot.dia, fck, self.fy)
            self._add("Beam", mark, level, prefix, "Bottom main", bot.dia, bot.count, cut, shape, dims, lap=lap)
        else:
            self.bbs.warnings.append(f"Beam {mark} ({level}): no bottom bars in the design – not scheduled")

        # top bars at supports
        top_ext = [[0.0, 0.0] for _ in segs]  # how far top bars reach into each segment from (start, end)
        top_dia: list[list[int]] = [[] for _ in segs]
        for j in range(n + 1):
            if sup[j][0] == "free":
                continue
            left = segs[j - 1] if j > 0 else None
            right = segs[j] if j < n else None
            bs = _larger(left.top_r_bars if left else None, right.top_l_bars if right else None)
            if bs is None:
                continue
            dia = bs.dia
            main, legs = 0.0, []
            for side, k in ((left, j - 1), (right, j)):
                if side is None:
                    e, lg = end_detail(j, dia)
                    main += e
                    legs.append(lg)
                    continue
                far = j - 1 if side is left else j + 1
                if sup[far][0] == "free":  # cantilever – run to the tip and bend down
                    main += spans[k] - c
                    legs.append(D - 2 * c)
                    ext = spans[k]
                else:
                    main += 0.3 * spans[k]
                    ext = 0.3 * spans[k]
                top_ext[k][0 if side is right else 1] = ext
                top_dia[k].append(dia)
            if n == 1:
                desc = "Top at left support" if j == 0 else "Top at right support"
            else:
                desc = "Top at left end support" if j == 0 else "Top at right end support" if j == n else ""
                desc = desc or f"Top over interior support {j}"
            shape, dims, cut = bent_bar(main, legs, dia)
            lap = lap_length(dia, fck, self.fy)
            self._add("Beam", mark, level, prefix, desc, dia, bs.count, cut, shape, dims, lap=lap)

        # hanger bars between the top bar cut-offs
        for k in range(n):
            if sup[k][0] == "free" or sup[k + 1][0] == "free":
                continue
            gap = spans[k] - sum(top_ext[k])
            if gap <= 0:
                continue
            hd = min([12] + top_dia[k])
            lap_c = lap_length(hd, fck, self.fy, compression=True)
            main = gap + sum(lap_c for e in top_ext[k] if e > 0)
            desc = "Hanger bars" if n == 1 else f"Hanger bars, span {k + 1}"
            shape, dims, cut = bent_bar(main, [], hd)
            self._add("Beam", mark, level, prefix, desc, hd, 2, cut, shape, dims, lap=lap_length(hd, fck, self.fy))

        # stirrups
        counts: dict[Links, int] = defaultdict(int)
        for seg, L in zip(segs, spans):
            if seg.links is None:
                self.bbs.warnings.append(f"Beam {mark} ({level}): shear design failed – stirrups not scheduled")
                continue
            counts[seg.links] += _ceil_count(L, seg.links.spacing)
        for lk, cnt in counts.items():
            self._add_links("Beam", mark, level, prefix, "Stirrup", lk, cnt, first.b * 1000, D, c)

    # ------------------------------------------------------------ columns
    def columns(self):
        foot = {f.mark: f for f in self.rep.footings}
        for cf in self.rep.combined_footings:  # only the depth D is used for the column foot
            for m in cf.marks:
                foot.setdefault(m, cf)
        c = self.ds.column_cover * 1000
        for cd in sorted(self.rep.columns, key=lambda x: (x.level_index, _natural(x.mark))):
            mem = self.fm.members.get(cd.member_id)
            if mem is None:
                continue
            prefix = f"{cd.mark}-L{cd.level_index}"
            fck = grade_fck(mem.grade)
            H = (cd.height or abs(self.fm.nodes[mem.n2].z - self.fm.nodes[mem.n1].z)) * 1000
            mb = cd.main_bars
            if mb is None or mb.count <= 0:
                self.bbs.warnings.append(f"Column {cd.mark} ({cd.level}): no main bars in the design – not scheduled")
            else:
                dia = mb.dia
                Ld = development_length(dia, fck, self.fy)
                lap = lap_length(dia, fck, self.fy)
                main, legs, notes = H, [], []
                if mem.n2 in self.col_above:
                    main += lap
                    notes.append("lap above floor")
                else:
                    cb = self.ds.beam_cover * 1000
                    Db = max((m.d for m in self.beams_at.get(mem.n2, [])), default=0.0) * 1000
                    main -= cb
                    legs.append(anchorage_leg(Ld, max(Db - cb, 0.0), dia) or 12 * dia)
                    notes.append("bent into roof beam")
                f = foot.get(cd.mark)
                if self.fm.nodes[mem.n1].support and f is not None:
                    cf = self.ds.footing_cover * 1000
                    inside = f.D * 1000 - cf
                    main += inside
                    legs.insert(0, anchorage_leg(Ld, inside, dia) or 12 * dia)
                    notes.append("foot in footing")
                shape, dims, cut = bent_bar(main, legs, dia)
                desc = "Main bar" + (f" ({', '.join(notes)})" if notes else "")
                self._add("Column", cd.mark, cd.level, prefix, desc, dia, mb.count, cut, shape, dims, lap=lap)
            if cd.tie is None:
                self.bbs.warnings.append(f"Column {cd.mark} ({cd.level}): no ties in the design – not scheduled")
            else:
                cnt = _ceil_count(H, cd.tie.spacing)
                self._add_links("Column", cd.mark, cd.level, prefix, "Tie", cd.tie, cnt, cd.b * 1000, cd.d * 1000, c)
        for w in self.rep.walls:
            self.bbs.warnings.append(
                f"Wall {w.mark} ({w.level}): shear wall bars not scheduled – see the wall design "
                f"(vertical {w.vertical}; horizontal {w.horizontal})"
            )

    # ------------------------------------------------------------ footings
    def footings(self):
        groups: dict[tuple, list] = defaultdict(list)
        for f in self.rep.footings:
            groups[(f.L, f.B, f.D, f.mesh_L, f.mesh_B)].append(f)
        c = self.ds.footing_cover * 1000
        fck = grade_fck(self.p.levels[0].grade) if self.p.levels else 25.0
        order = sorted(groups.values(), key=lambda fs: _natural(fs[0].mark))
        for i, fs in enumerate(order, start=1):
            f = fs[0]
            member = ", ".join(sorted((x.mark for x in fs), key=_natural))
            prefix = f"F{i}"
            L, B, D = f.L * 1000, f.B * 1000, f.D * 1000
            for mesh, along, across, name in ((f.mesh_L, L, B, "L"), (f.mesh_B, B, L, "B")):
                if mesh is None:
                    self.bbs.warnings.append(f"Footing {member}: no bars along {name} in the design")
                    continue
                cnt = _ceil_count(across - 2 * c, mesh.spacing)
                shape, dims, cut = bent_bar(along - 2 * c, [D - 2 * c, D - 2 * c], mesh.dia)
                lap = lap_length(mesh.dia, fck, self.fy)
                desc = f"Bottom bars along {name} (ends bent up)"
                self._add("Footing", member, "Foundation", prefix, desc, mesh.dia, cnt, cut, shape, dims, len(fs), lap)
        for cf in self.rep.combined_footings:
            self.bbs.warnings.append(
                f"Combined footing {' + '.join(cf.marks)}: bars not scheduled – see the design "
                f"(bottom {cf.bottom}; top {cf.top})"
            )

    # ------------------------------------------------------------ slabs
    def slabs(self):
        by_plan: dict[str, list] = defaultdict(list)
        for pname, res in self.rep.slabs:
            by_plan[pname].append(res)
        c = self.ds.slab_cover * 1000
        for pname, results in by_plan.items():
            plan = self.p.plan(pname)
            levels = [(i, lv) for i, lv in enumerate(self.p.levels, start=1) if lv.plan == pname]
            strips = self._slab_strips(plan, results) if plan else []
            for li, lv in levels:
                fck = grade_fck(lv.grade)
                for s in sorted(results, key=lambda r: _natural(r.mark)):
                    self._slab(s, lv.name, f"{s.mark}-L{li}", fck, c)
                for a, b, overlap in strips:
                    self._strip(a, b, overlap, lv.name, li, fck)

    def _slab(self, s, level: str, prefix: str, fck: float, c: float):
        lx, ly, D = s.lx * 1000, s.ly * 1000, s.D_mm

        def add(desc, mesh: BarMesh | None, length, across, legs=()):
            if mesh is None:
                return
            shape, dims, cut = bent_bar(length, list(legs), mesh.dia)
            lap = lap_length(mesh.dia, fck, self.fy)
            cnt = _ceil_count(across, mesh.spacing)
            self._add("Slab", s.mark, level, prefix, desc, mesh.dia, cnt, cut, shape, dims, lap=lap)

        if s.kind == "cantilever":
            main = _larger(s.mesh_neg, s.mesh_x)
            if main is not None:
                Ld = development_length(main.dia, fck, self.fy)
                add("Top main (cantilever, into back span)", main, lx - c + Ld, ly, [D - 2 * c])
            add("Top distribution", s.mesh_y, ly, lx)
            return
        one_way = s.kind == "one_way"
        add("Bottom main, short span" if one_way else "Bottom, short span", s.mesh_x, lx, ly)
        add("Distribution, long span" if one_way else "Bottom, long span", s.mesh_y, ly, lx)

    def _slab_strips(self, plan, results) -> list[tuple]:
        """Continuous edges shared by two designed (non-cantilever) panels: (res_a, res_b, overlap m)."""
        res = {r.mark: r for r in results}
        slabs = [s for s in plan.slabs if s.mark in res and s.distribution != "on_grade" and len(s.points) >= 3]
        out = []
        for i, sa in enumerate(slabs):
            for sb in slabs[i + 1 :]:
                ra, rb = res[sa.mark], res[sb.mark]
                if sa.mark == sb.mark or "cantilever" in (ra.kind, rb.kind):
                    continue
                if ra.mesh_neg is None and rb.mesh_neg is None:
                    continue
                overlap = 0.0
                for a, b in sa.edges():
                    for c_, d_ in sb.edges():
                        hit = G.collinear_overlap(a, b, c_, d_, tol=0.02)
                        if hit:
                            overlap = max(overlap, (hit[1] - hit[0]) * G.dist(a, b))
                if overlap > 0.05:
                    out.append((ra, rb, overlap))
        return out

    def _strip(self, a, b, overlap: float, level: str, li: int, fck: float):
        mesh = _larger(a.mesh_neg, b.mesh_neg)
        member = f"{a.mark}/{b.mark}"
        prefix = f"{a.mark}-{b.mark}-L{li}"
        width = 0.25 * (a.lx + b.lx) * 1000
        lap = lap_length(mesh.dia, fck, self.fy)
        cnt = _ceil_count(overlap * 1000, mesh.spacing)
        desc = "Top over continuous edge (0.25 lx each side)"
        self._add("Slab", member, level, prefix, desc, mesh.dia, cnt, width, "00", {"a": width}, lap=lap)
        dist = _larger(a.mesh_y, b.mesh_y)
        if dist is not None:
            cnt = _ceil_count(width, dist.spacing)
            L = overlap * 1000
            self._add("Slab", member, level, prefix, "Distribution under top bars", dist.dia, cnt, L, "00", {"a": L})


def build_bbs(project: Project, fm: FrameModel, rep: DesignReport) -> BBS:
    """Bar bending schedule for every designed beam, column, footing and slab (see module docs)."""
    b = _Builder(project, fm, rep)
    b.beams()
    b.columns()
    b.footings()
    b.slabs()
    return b.bbs


# ---------------------------------------------------------------- Excel
HEADERS = [
    "Member",
    "Level",
    "Bar mark",
    "Description",
    "Shape code",
    "Dia (mm)",
    "No./member",
    "Members",
    "Total no.",
    "Cutting length (m)",
    "Total length (m)",
    "Unit wt (kg/m)",
    "Weight (kg)",
    "a (mm)",
    "b (mm)",
    "c (mm)",
]


def write_bbs_excel(path: str, bbs: BBS, project: Project, watermark: str = "") -> str:
    """Write the schedule to an Excel workbook: Summary + one sheet per member type."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    note_font = Font(italic=True, color="9C0006")

    def style_header(ws, row: int):
        for cell in ws[row]:
            if cell.value is not None:
                cell.fill, cell.font = head_fill, head_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def footer(ws):
        ws.append([])
        if watermark:
            ws.append([watermark])
        ws.append([CUT_NOTE])
        ws.cell(ws.max_row, 1).font = note_font
        ws.append([DISCLAIMER])

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"{APP_NAME} {__version__} – Bar Bending Schedule"])
    ws["A1"].font = Font(size=14, bold=True)
    info = [
        ("Project", project.name),
        ("Client", project.client),
        ("Engineer", project.engineer),
        ("Location", project.location),
        ("Date", _dt.date.today().isoformat()),
        ("Prepared by", COMPANY),
        ("Codes", "IS 456:2000, IS 2502:1963, SP 34:1987"),
        ("Steel grade", f"Fe {project.design.fy_main:.0f}"),
    ]
    if watermark:
        info.append(("Licence", watermark))
    for k, v in info:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    total = bbs.total_weight
    totals = [
        ("Total steel (kg)", round(total, 1)),
        ("Total steel (t)", round(total / 1000, 3)),
    ]
    if bbs.concrete_m3 > 0:
        totals += [
            ("Structural concrete (m³)", round(bbs.concrete_m3, 2)),
            ("Steel / concrete (kg/m³)", round(bbs.steel_per_m3, 1)),
        ]
    if bbs.boq_steel > 0:
        totals.append(("Design BOQ steel estimate (kg)", round(bbs.boq_steel, 1)))
    for k, v in totals:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold

    def table(title, key_head, data: dict):
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = bold
        ws.append([key_head, "Weight (kg)", "Weight (t)", "Share %"])
        style_header(ws, ws.max_row)
        for k, v in data.items():
            ws.append([k, round(v, 1), round(v / 1000, 3), round(100 * v / total, 1) if total else 0.0])
        ws.append(["Total", round(total, 1), round(total / 1000, 3), 100.0 if total else 0.0])
        for cell in ws[ws.max_row]:
            cell.font = bold

    table("By diameter", "Diameter", {f"T{d}": w for d, w in bbs.by_dia().items()})
    table("By level", "Level", bbs.by_level())
    by_type = bbs.by_member_type()
    table("By member type", "Member type", {t: by_type[t] for t in MEMBER_TYPES if t in by_type})
    ws.append([])
    ws.append(["Shape code", "Description"])
    style_header(ws, ws.max_row)
    for code, desc in SHAPE_CODES.items():
        ws.append([code, desc])
    if bbs.warnings:
        ws.append([])
        ws.append(["Warnings"])
        ws.cell(ws.max_row, 1).font = bold
        for w in bbs.warnings:
            ws.append([w])
    footer(ws)
    ws.column_dimensions["A"].width = 34
    for col in "BCD":
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["B"].width = 48

    # ---- member sheets
    sheet_names = {"Beam": "Beams", "Column": "Columns", "Footing": "Footings", "Slab": "Slabs"}
    for mtype in MEMBER_TYPES:
        ws = wb.create_sheet(sheet_names[mtype])
        ws.append(HEADERS)
        style_header(ws, 1)
        items = [it for it in bbs.items if it.member_type == mtype]
        rows = []
        for it in items:
            rows.append(
                [
                    it.member,
                    it.level,
                    it.bar_mark,
                    it.description,
                    it.shape_code,
                    it.dia,
                    it.count,
                    it.members,
                    it.total_count,
                    round(it.cutting_length, 3),
                    round(it.total_length, 2),
                    round(it.unit_weight, 3),
                    round(it.weight, 2),
                    it.dims.get("a"),
                    it.dims.get("b"),
                    it.dims.get("c"),
                ]
            )
            ws.append(rows[-1])
        ws.append(["TOTAL"] + [""] * 9 + [round(sum(it.total_length for it in items), 2), "", None])
        ws.cell(ws.max_row, 13).value = round(sum(it.weight for it in items), 2)
        for cell in ws[ws.max_row]:
            cell.font = bold
        for i, h in enumerate(HEADERS, 1):
            width = max([len(h)] + [len(str(r[i - 1])) for r in rows[:300]] + [6])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 48)
        ws.freeze_panes = "A2"
        footer(ws)
    literal_text(wb)  # names/marks starting with '=' stay text
    wb.save(path)
    return path
