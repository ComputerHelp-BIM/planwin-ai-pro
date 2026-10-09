"""Results navigator: a grouped tree of result tables with a row filter, copy and Excel export.

The left tree lists every table under its group (Plan, Analysis, Design, Code checks,
Quantities) with a badge – the number of issues, the number of failing members (✖ n) or
a tick when every row passes.  The right side shows the selected table with a filter box,
a row count and Copy / Excel buttons.  The export helpers are plain functions so they can
be used and tested without dialogs.
"""

from __future__ import annotations

import os
import re
from typing import TYPE_CHECKING

from PySide6.QtCore import Qt, Signal
from PySide6.QtGui import QColor, QKeySequence, QShortcut
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMenu,
    QMessageBox,
    QSplitter,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QToolButton,
    QTreeWidget,
    QTreeWidgetItem,
    QVBoxLayout,
    QWidget,
)

from .. import units
from ..io.report_common import literal_text
from ..services.exports import safe_filename
from . import theme

if TYPE_CHECKING:
    from .main_window import MainWindow

#: every results table, in creation order (``ResultsPanel.tables`` keeps this order)
RESULT_TABS = (
    "Issues",
    "Column loads",
    "Beam loads",
    "Columns",
    "Beams",
    "Footings",
    "Slabs",
    "Lateral",
    "Drift",
    "BOQ & cost",
    "Walls",
    "Combined footings",
    "IS 13920",
    "Irregularity",
    "Modal / RS",
    "BOQ by floor",
)
#: tables filled from the design report (emptied when the results are invalidated)
DESIGN_TABS = (
    "Columns",
    "Beams",
    "Footings",
    "Slabs",
    "Drift",
    "BOQ & cost",
    "Walls",
    "Combined footings",
    "IS 13920",
    "Irregularity",
    "Modal / RS",
    "BOQ by floor",
)
#: navigator groups, in display order
RESULT_GROUPS = (
    ("Plan", ("Issues", "Column loads", "Beam loads")),
    ("Analysis", ("Lateral", "Drift", "Modal / RS")),
    ("Design", ("Columns", "Beams", "Walls", "Footings", "Combined footings", "Slabs")),
    ("Code checks", ("IS 13920", "Irregularity")),
    ("Quantities", ("BOQ & cost", "BOQ by floor")),
)
#: one-line description of every table (navigator tooltip)
TABLE_INFO = {
    "Issues": "Issues: plan and frame checks – unsupported beams, missing columns, open slabs, model warnings",
    "Column loads": "Column loads: slab → beam → column load take-down of the current plan (dead, live, total)",
    "Beam loads": "Beam loads: span, total dead / live load, equivalent UDL and supports of each beam",
    "Lateral": "Lateral: IS 1893 seismic (T, Sa/g, Ah, W, base shear) and IS 875-3 wind base shear per case",
    "Drift": "Drift: storey drift per load case and level against the IS 1893 limit of 0.004 h",
    "Modal / RS": "Modal / RS: mode periods, mass participation and response spectrum base shear scaling",
    "Columns": "Columns: IS 456 biaxial design per level – size, Pu, Mux/Muy, steel, ties, utilisation",
    "Beams": "Beams: IS 456 flexure and shear design per level – moments, shear, bars and stirrups",
    "Walls": "Walls: IS 456 / IS 13920 shear wall design – axial, moment, shear, curtains and boundary elements",
    "Footings": "Footings: isolated footing size, depth, bars and soil pressure against SBC",
    "Combined footings": "Combined footings: two-column footings – size, depth, soil pressure and reinforcement",
    "Slabs": "Slabs: IS 456 one-way / two-way slab design – depth, short / long / top steel",
    "IS 13920": "IS 13920: ductile detailing clause checks for beams, columns, joints and walls",
    "Irregularity": "Irregularity: IS 1893 Table 5 / 6 plan and vertical irregularity checks",
    "BOQ & cost": "BOQ & cost: concrete, steel and formwork quantities with rates and amounts",
    "BOQ by floor": "BOQ by floor: concrete, steel, formwork and cost by floor and by member type",
}
#: tables whose row order matters (sorting stays off)
_UNSORTED = ("Issues", "Modal / RS", "BOQ by floor")
_BAD = ("NO", False, "error")
_DEFAULT_BADGE = {"error": "#DC2626", "warn": "#D97706", "ok": "#16A34A"}


# ---------------------------------------------------------------------- export helpers
def _cell_value(item: QTableWidgetItem | None):
    """The value of a cell as shown: numbers stay numbers, everything else is text."""
    if item is None:
        return ""
    v = item.data(Qt.DisplayRole)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    return "" if v is None else str(v)


def table_rows(t: QTableWidget, visible_only: bool = True, rows=None) -> tuple[list[str], list[list]]:
    """Headers and cell values of a table in display order (sorted, filtered).

    ``visible_only`` skips rows hidden by the filter; ``rows`` limits the result to those row
    indices (e.g. the selection)."""
    headers = []
    for c in range(t.columnCount()):
        h = t.horizontalHeaderItem(c)
        headers.append(h.text() if h else "")
    wanted = set(rows) if rows is not None else None
    out = []
    for r in range(t.rowCount()):
        if (visible_only and t.isRowHidden(r)) or (wanted is not None and r not in wanted):
            continue
        out.append([_cell_value(t.item(r, c)) for c in range(t.columnCount())])
    return headers, out


def rows_tsv(headers: list[str], rows: list[list]) -> str:
    """Tab separated text with a header line (pastes into Excel as cells)."""

    def cell(v) -> str:
        return re.sub(r"[\t\r\n]+", " ", str(v))

    return "\n".join("\t".join(cell(v) for v in line) for line in [headers, *rows]) + "\n"


def sheet_title(name: str, used: set[str] | None = None) -> str:
    """A valid, unique Excel sheet name (≤ 31 characters, no ``[]:*?/\\``)."""
    base = re.sub(r"[\[\]:*?/\\]", "_", name).strip("'") or "Sheet"
    title = base[:31]
    used = used if used is not None else set()
    k = 2
    while title.lower() in {u.lower() for u in used}:
        sfx = f" ({k})"
        title = base[: 31 - len(sfx)] + sfx
        k += 1
    used.add(title)
    return title


def write_tables_xlsx(path: str, tables: dict[str, tuple[list[str], list[list]]]) -> str:
    """Write one sheet per table (bold, frozen header row; numbers stay numeric)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    bold = Font(bold=True)
    for name, (headers, rows) in tables.items():
        ws = wb.create_sheet(sheet_title(name, used))
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = bold
        ws.freeze_panes = "A2"
        widths = [len(str(h)) for h in headers]
        for row in rows:
            ws.append(list(row))
            for c, v in enumerate(row):
                if c < len(widths):
                    widths[c] = max(widths[c], len(str(v)))
                else:
                    widths.append(len(str(v)))
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 6), 60)
    if not wb.sheetnames:
        wb.create_sheet("Results")
    literal_text(wb)  # cell text such as "=B1" stays text
    wb.save(path)
    return path


# ---------------------------------------------------------------------- navigator
class _NavTree(QTreeWidget):
    """Tree whose Up / Down keys step over the group rows straight to the next table."""

    def keyPressEvent(self, ev):
        if ev.key() in (Qt.Key_Up, Qt.Key_Down) and not ev.modifiers():
            order = []
            for g in range(self.topLevelItemCount()):
                grp = self.topLevelItem(g)
                order.append(grp)
                if grp.isExpanded():
                    order += [grp.child(i) for i in range(grp.childCount())]
            cur = self.currentItem()
            i = order.index(cur) if cur in order else -1
            step = 1 if ev.key() == Qt.Key_Down else -1
            i += step
            while 0 <= i < len(order):
                if order[i].parent() is not None:
                    self.setCurrentItem(order[i])
                    break
                i += step
            ev.accept()
            return
        super().keyPressEvent(ev)


class ResultsPanel(QWidget):
    """Grouped results navigator (tree on the left, the selected table on the right)."""

    issueActivated = Signal(object)

    def __init__(self, main: MainWindow):
        super().__init__()
        self.main = main
        self.tables: dict[str, QTableWidget] = {}
        self._issues: list = []
        self._items: dict[str, QTreeWidgetItem] = {}
        self._bad: dict[str, int | None] = {}  # failing rows per table (None: no pass / fail column)
        self._have_checks = False
        self._current = RESULT_TABS[0]
        self._restoring = True

        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        self.splitter = QSplitter(Qt.Horizontal)
        self.splitter.setChildrenCollapsible(False)
        lay.addWidget(self.splitter)

        # ---- left: navigator
        self.nav = _NavTree()
        self.nav.setObjectName("ResultsNav")
        self.nav.setHeaderHidden(True)
        self.nav.setColumnCount(2)
        self.nav.setIndentation(0)  # the stylesheet indents child rows (an indent area paints a 2nd selection)
        self.nav.setRootIsDecorated(False)
        self.nav.setSelectionMode(QAbstractItemView.SingleSelection)
        self.nav.setAllColumnsShowFocus(True)
        self.nav.setMinimumWidth(130)
        hdr = self.nav.header()
        hdr.setStretchLastSection(False)
        hdr.setSectionResizeMode(0, QHeaderView.Stretch)
        hdr.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        for group, names in RESULT_GROUPS:
            g = QTreeWidgetItem(self.nav, [group, ""])
            g.setFlags(Qt.ItemIsEnabled)
            f = g.font(0)
            f.setBold(True)
            g.setFont(0, f)
            g.setFirstColumnSpanned(True)
            for name in names:
                it = QTreeWidgetItem(g, [name, ""])
                it.setData(0, Qt.UserRole, name)
                it.setToolTip(0, TABLE_INFO.get(name, name))
                it.setToolTip(1, TABLE_INFO.get(name, name))
                it.setTextAlignment(1, Qt.AlignRight | Qt.AlignVCenter)
                self._items[name] = it
            g.setExpanded(True)
        self.nav.currentItemChanged.connect(self._nav_changed)
        self.nav.itemClicked.connect(self._nav_clicked)
        self.splitter.addWidget(self.nav)

        # ---- right: toolbar + tables
        right = QWidget()
        rv = QVBoxLayout(right)
        rv.setContentsMargins(6, 2, 4, 0)
        rv.setSpacing(4)
        bar = QHBoxLayout()
        bar.setSpacing(6)
        self.filter = QLineEdit()
        self.filter.setObjectName("ResultsFilter")
        self.filter.setPlaceholderText("Filter rows…")
        self.filter.setClearButtonEnabled(True)
        self.filter.setToolTip("Show only the rows containing this text (any column, Ctrl+F)")
        self.filter.textChanged.connect(self._apply_filter)
        bar.addWidget(self.filter, 1)
        self.count_lbl = QLabel()
        self.count_lbl.setObjectName("ResultsCount")
        bar.addWidget(self.count_lbl)
        self.copy_btn = QToolButton()
        self.copy_btn.setObjectName("TableAction")
        self.copy_btn.setText("Copy")
        self.copy_btn.setToolTip("Copy the selected rows (or the whole visible table) with headers – pastes into Excel")
        self.copy_btn.clicked.connect(self.copy_rows)
        bar.addWidget(self.copy_btn)
        self.excel_btn = QToolButton()
        self.excel_btn.setObjectName("TableAction")
        self.excel_btn.setText("Excel")
        self.excel_btn.setToolTip("Export the visible rows of this table to an Excel workbook")
        self.excel_btn.setPopupMode(QToolButton.MenuButtonPopup)
        menu = QMenu(self.excel_btn)
        menu.addAction("Excel (this table)", self.export_current)
        menu.addAction("Excel (all tables)", self.export_all)
        self.excel_btn.setMenu(menu)
        self.excel_btn.clicked.connect(self.export_current)
        bar.addWidget(self.excel_btn)
        rv.addLayout(bar)

        self.stack = QStackedWidget()
        for name in RESULT_TABS:
            t = QTableWidget()
            t.setEditTriggers(QAbstractItemView.NoEditTriggers)
            t.setSelectionBehavior(QAbstractItemView.SelectRows)
            t.setAlternatingRowColors(True)
            t.verticalHeader().setVisible(False)
            t.setSortingEnabled(True)
            self.tables[name] = t
            self.stack.addWidget(t)
        self.tables["Issues"].cellDoubleClicked.connect(self._issue_clicked)
        self.empty = QLabel()
        self.empty.setObjectName("ResultsEmpty")
        self.empty.setAlignment(Qt.AlignCenter)
        self.empty.setWordWrap(True)
        self.area = QStackedWidget()
        self.area.addWidget(self.stack)
        self.area.addWidget(self.empty)
        rv.addWidget(self.area, 1)
        self.splitter.addWidget(right)
        self.splitter.setStretchFactor(0, 0)
        self.splitter.setStretchFactor(1, 1)
        self.splitter.setSizes([205, 900])
        self.splitter.splitterMoved.connect(self._save_splitter)

        sc = QShortcut(QKeySequence.Find, self, self._focus_filter)
        sc.setContext(Qt.WidgetWithChildrenShortcut)

        for name in RESULT_TABS:
            self._bad[name] = None
            self._update_item(name)
        self._restore()
        self._restoring = False

    # ------------------------------------------------------------------ navigation
    def show_tab(self, name: str) -> None:
        """Show the results table ``name`` (one of :data:`RESULT_TABS`) and select it in the navigator."""
        it = self._items[name]
        it.parent().setExpanded(True)
        self.nav.setCurrentItem(it)
        self._show(name)

    def current_name(self) -> str:
        """Name of the table on show."""
        return self._current

    # the former QTabWidget API, kept for callers that still use it
    def currentIndex(self) -> int:
        return RESULT_TABS.index(self._current)

    def tabText(self, index: int) -> str:
        return RESULT_TABS[index]

    def _nav_changed(self, cur, _prev):
        if cur is not None and cur.parent() is not None:
            self._show(cur.data(0, Qt.UserRole))

    def _nav_clicked(self, item, _col):
        if item.parent() is None:  # a group: toggle it, keep the table selection
            item.setExpanded(not item.isExpanded())
            self.nav.setCurrentItem(self._items[self._current])

    def _show(self, name: str) -> None:
        self._current = name
        self.stack.setCurrentWidget(self.tables[name])
        self._apply_filter()
        s = self._settings()
        if s is not None and not self._restoring:
            s.setValue("results/table", name)

    def _focus_filter(self):
        self.filter.setFocus(Qt.ShortcutFocusReason)
        self.filter.selectAll()

    # ------------------------------------------------------------------ settings
    def _settings(self):
        return getattr(self.main, "settings", None)

    def _restore(self):
        s = self._settings()
        name = RESULT_TABS[0]
        if s is not None:
            try:
                saved = s.value("results/table", name)
                if saved in self.tables:
                    name = saved
                sizes = [int(x) for x in (s.value("results/splitter") or [])]
                if len(sizes) == 2 and min(sizes) > 0:
                    self.splitter.setSizes(sizes)
            except (TypeError, ValueError):
                pass
        self.show_tab(name)

    def _save_splitter(self, *_):
        s = self._settings()
        if s is not None:
            s.setValue("results/splitter", self.splitter.sizes())

    # ------------------------------------------------------------------ filter / state
    def _apply_filter(self, *_):
        t = self.tables[self._current]
        text = self.filter.text().strip().lower()
        shown = 0
        for r in range(t.rowCount()):
            hit = not text or any(
                (it := t.item(r, c)) is not None and text in it.text().lower() for c in range(t.columnCount())
            )
            t.setRowHidden(r, not hit)
            shown += hit
        n = t.rowCount()
        self.count_lbl.setText(f"{shown} of {n} rows" if text else f"{n} row{'s' * (n != 1)}")
        if n:
            self.area.setCurrentWidget(self.stack)
        else:
            self.empty.setText(self._empty_text(self._current))
            self.area.setCurrentWidget(self.empty)
        self.copy_btn.setEnabled(n > 0)

    def _empty_text(self, name: str) -> str:
        filled = self.tables[name].columnCount() > 0
        if name == "Issues":
            if self._have_checks:
                return "No issues – the plan and frame checks found nothing to report."
            return "Run Analyse plan (F5) or Analyse frame (F6) to check the model."
        if name in ("Column loads", "Beam loads"):
            return "Run Analyse plan (F5) to see the load take-down."
        if name in DESIGN_TABS:
            if filled:
                return f"No rows in {name} for this design."
            return "Run Design all (F7) to see results."
        if filled:
            return f"No rows in {name}."
        return "Run Analyse frame (F6) to see results."

    def _palette(self) -> dict:
        pal = getattr(theme, "PALETTES", {})
        return pal.get(getattr(self.main, "theme_name", "light"), pal.get("light", {}))

    def _update_item(self, name: str) -> None:
        it = self._items[name]
        t = self.tables[name]
        n = t.rowCount()
        badge = getattr(theme, "BADGE", _DEFAULT_BADGE)
        bad = self._bad.get(name)
        text, color, tip = "", None, TABLE_INFO.get(name, name)
        if name == "Issues":
            if n:
                text, color = f"({n})", badge["error"] if bad else badge["warn"]
                tip = f"{n} issue{'s' * (n != 1)}" + (f", {bad} error{'s' * (bad != 1)}" if bad else "")
        elif bad is not None and n:
            if bad:
                text, color, tip = f"✖ {bad}", badge["error"], f"{bad} of {n} rows fail"
            else:
                text, color, tip = "✓", badge["ok"], f"all {n} rows pass"
        it.setText(1, text)
        it.setToolTip(1, tip)
        it.setData(1, Qt.ForegroundRole, QColor(color) if color else None)
        muted = self._palette().get("muted")
        it.setData(0, Qt.ForegroundRole, QColor(muted) if n == 0 and muted else None)
        f = it.font(1)
        f.setBold(bool(text))
        it.setFont(1, f)

    def restyle(self) -> None:
        """Recolour the navigator after a theme switch (its colours are set per item, not by the style sheet)."""
        for name in self._items:
            self._update_item(name)

    def _after_change(self, name: str) -> None:
        self._update_item(name)
        if name == self._current:
            self._apply_filter()

    # ------------------------------------------------------------------ copy / export
    def _status(self, msg: str) -> None:
        sb = getattr(self.main, "statusBar", None)
        if callable(sb):
            sb().showMessage(msg, 5000)

    def copy_rows(self) -> str:
        """Copy the selected rows – or every visible row – of the current table as TSV with headers."""
        t = self.tables[self._current]
        sel = {i.row() for i in t.selectionModel().selectedIndexes() if not t.isRowHidden(i.row())}
        headers, rows = table_rows(t, rows=sel or None)
        text = rows_tsv(headers, rows) if headers else ""
        QApplication.clipboard().setText(text)
        self._status(f"Copied {len(rows)} row{'s' * (len(rows) != 1)} of {self._current}")
        return text

    def _ask_path(self, base: str) -> str:
        s = self._settings()
        last = os.path.expanduser("~")
        if s is not None:
            last = str(s.value("last_dir", last) or last)
        proj = getattr(getattr(self.main, "project", None), "name", "") or "project"
        default = os.path.join(last, f"{safe_filename(proj)}_{safe_filename(base, 'results')}.xlsx")
        fn, _ = QFileDialog.getSaveFileName(self, "Export to Excel", default, "Excel workbook (*.xlsx)")
        if fn and not fn.lower().endswith(".xlsx"):
            fn += ".xlsx"
        return fn

    def _write(self, base: str, tables: dict) -> str:
        if not tables:
            self._status("Nothing to export – the table is empty")
            return ""
        fn = self._ask_path(base)
        if not fn:
            return ""
        try:
            write_tables_xlsx(fn, tables)
        except OSError as exc:  # e.g. the workbook is open in Excel
            QMessageBox.warning(self, "Export to Excel", f"Could not write {fn}:\n{exc}")
            return ""
        s = self._settings()
        if s is not None:
            s.setValue("last_dir", os.path.dirname(fn))
        what = next(iter(tables)) if len(tables) == 1 else f"{len(tables)} tables"
        self._status(f"Exported {what} to {fn}")
        return fn

    def export_current(self) -> str:
        """Export the visible rows of the current table to an .xlsx chosen by the user."""
        headers, rows = table_rows(self.tables[self._current])
        return self._write(self._current, {self._current: (headers, rows)} if headers else {})

    def export_all(self) -> str:
        """Export every non-empty table (all rows) to one workbook, one sheet per table."""
        tables = {n: table_rows(t, visible_only=False) for n, t in self.tables.items() if t.rowCount()}
        return self._write("results", tables)

    # ------------------------------------------------------------------ filling
    def _fill(self, name: str, headers: list[str], rows: list[list], bad_col: int | None = None):
        """Fill a table. Numbers in a column whose header has a kN unit are shown in ``units.current``
        (force, moment, kN/m and kN/m² all convert by the same kN → t factor)."""
        u = units.current
        conv = [u.mks and "kN" in h for h in headers]
        t = self.tables[name]
        t.setSortingEnabled(False)
        t.clear()
        t.setRowCount(0)  # also drops the rows hidden by the filter
        t.setColumnCount(len(headers))
        t.setHorizontalHeaderLabels([u.text(h) for h in headers])
        t.setRowCount(len(rows))
        nbad = 0
        for r, row in enumerate(rows):
            bad = bad_col is not None and row[bad_col] in _BAD
            nbad += bad
            for c, v in enumerate(row):
                it = QTableWidgetItem()
                if isinstance(v, float):  # plain float(): numpy scalars are stored as opaque objects and show blank
                    it.setData(Qt.DisplayRole, round(float(u.show(v, "force") if conv[c] else v), 3))
                elif isinstance(v, int) and not isinstance(v, bool) and conv[c]:
                    it.setData(Qt.DisplayRole, round(u.show(float(v), "force"), 3))
                else:
                    it.setData(Qt.DisplayRole, int(v) if isinstance(v, int) and not isinstance(v, bool) else str(v))
                if bad:
                    it.setBackground(QColor(220, 38, 38, 60))
                t.setItem(r, c, it)
        t.resizeColumnsToContents()
        t.setSortingEnabled(name not in _UNSORTED)  # row order matters there
        self._bad[name] = nbad if bad_col is not None else None
        self._after_change(name)

    def _clear(self, *names: str):
        for n in names:
            t = self.tables[n]
            t.setSortingEnabled(False)
            t.clear()
            t.setRowCount(0)
            t.setColumnCount(0)
            self._bad[n] = None
            self._after_change(n)

    def refresh(self):
        """Rebuild every table. Tables whose results were invalidated by an edit are
        emptied so that out-of-date design results are never shown next to a changed model."""
        m = self.main
        plan = m.current_plan()
        issues = []
        self._issues = []
        res = m.plan_result(plan.name) if plan else None
        if not res:
            self._clear("Column loads", "Beam loads")
        if plan:
            if res:
                for i in res.issues:
                    issues.append([plan.name, i.level, i.kind, i.message])
                    self._issues.append(i)
                self._fill(
                    "Column loads",
                    ["Column", "Dead kN", "Live kN", "Total kN", "From beams"],
                    [[c.mark, c.dead, c.live, c.total, ", ".join(p[0] for p in c.parts)] for c in res.columns.values()],
                )
                bmap = {b.id: b for b in plan.beams}
                rows = []
                for bid, br in res.beams.items():
                    b = bmap.get(bid)
                    if b:
                        rows.append(
                            [
                                b.mark,
                                br.length,
                                br.total("D"),
                                br.total("L"),
                                br.equivalent_udl(),
                                " / ".join(f"{s.kind[0].upper()}@{s.x:.2f}" for s in br.supports),
                            ]
                        )
                self._fill(
                    "Beam loads", ["Beam", "Span m", "Total D kN", "Total L kN", "Eq. UDL kN/m", "Supports"], rows
                )
        fm = m.frame_model()
        if not fm:
            self._clear("Lateral")
        if fm:
            for i in fm.issues:
                issues.append(["Frame", i.level, i.kind, i.message])
                self._issues.append(i)
            rows = [[k, "Seismic", r.T, r.sa_g, r.Ah, r.W, r.Vb] for k, r in fm.seismic.items()]
            rows += [[k, "Wind", "", "", "", "", sum(r.forces)] for k, r in fm.wind.items()]
            self._fill("Lateral", ["Case", "Type", "T (s)", "Sa/g", "Ah", "W (kN)", "Base shear (kN)"], rows)
        self._have_checks = bool(res or fm)
        self._fill("Issues", ["Where", "Level", "Type", "Message"], issues, bad_col=1)
        rep = m.design_report()
        if not rep:
            self._clear(*DESIGN_TABS)
            return
        self._fill_members(rep)
        self._fill_v110(rep, m.frame_analysis(), fm)

    def _fill_members(self, rep):
        self._fill(
            "Columns",
            [
                "Level",
                "Col",
                "b",
                "D",
                "Pu kN",
                "Mux kN·m",
                "Muy kN·m",
                "p %",
                "Bars",
                "Ties",
                "Ratio",
                "OK",
                "Governing",
            ],
            [
                [
                    c.level,
                    c.mark,
                    c.b,
                    c.d,
                    c.Pu,
                    c.Mux,
                    c.Muy,
                    c.steel_pct,
                    c.bars,
                    c.ties,
                    c.utilisation,
                    "YES" if c.ok else "NO",
                    c.governing,
                ]
                for c in rep.columns
            ],
            bad_col=11,
        )
        self._fill(
            "Beams",
            [
                "Level",
                "Beam",
                "b",
                "D",
                "Span",
                "Mu+ kN·m",
                "Mu- L kN·m",
                "Mu- R kN·m",
                "Vu kN",
                "Bottom",
                "Top L",
                "Top R",
                "Stirrups",
                "OK",
                "Notes",
            ],
            [
                [
                    b.level,
                    b.mark,
                    b.b,
                    b.d,
                    b.span,
                    b.M_sag,
                    b.M_hog_l,
                    b.M_hog_r,
                    b.V_max,
                    b.bottom,
                    b.top_l,
                    b.top_r,
                    b.stirrups,
                    "YES" if b.ok else "NO",
                    "; ".join(b.notes),
                ]
                for b in rep.beams
            ],
            bad_col=13,
        )
        self._fill(
            "Footings",
            ["Col", "P kN", "L m", "B m", "D m", "Bars ∥L", "Bars ∥B", "q kN/m²", "OK"],
            [
                [f.mark, f.P_service, f.L, f.B, f.D, f.bars_L, f.bars_B, f.q, "YES" if f.ok else "NO"]
                for f in rep.footings
            ],
            bad_col=8,
        )
        self._fill(
            "Slabs",
            ["Plan", "Slab", "lx", "ly", "Type", "D mm", "Short", "Long", "Top", "OK", "Notes"],
            [
                [
                    pn,
                    s.mark,
                    s.lx,
                    s.ly,
                    s.kind,
                    s.D_mm,
                    s.ast_x,
                    s.ast_y,
                    s.ast_neg,
                    "YES" if s.ok else "NO",
                    "; ".join(s.notes),
                ]
                for pn, s in rep.slabs
            ],
            bad_col=9,
        )
        self._fill(
            "Drift",
            ["Case", "Level", "Drift mm", "Ratio", "OK"],
            [[d["case"], d["level"], d["drift_mm"], d["ratio"], "YES" if d["ok"] else "NO"] for d in rep.drifts],
            bad_col=4,
        )
        lines = [list(x) for x in rep.boq.get("lines", [])]
        lines.append(["TOTAL", "", "", "", rep.boq.get("cost", 0.0)])
        lines.append(["Steel/concrete ratio", "kg/m³", rep.boq.get("steel_per_m3", 0.0), "", ""])
        self._fill("BOQ & cost", ["Item", "Unit", "Quantity", "Rate ₹", "Amount ₹"], lines)

    def _fill_v110(self, rep, fa, fm):
        """Walls, combined footings, IS 13920, irregularity, modal / response spectrum and floor-wise BOQ."""
        yes = {True: "YES", False: "NO"}
        self._fill(
            "Walls",
            [
                "Level",
                "Wall",
                "t m",
                "Lw m",
                "Pu kN",
                "Mu kN·m",
                "Vu kN",
                "ρv %",
                "ρh %",
                "Curtains",
                "Vertical",
                "Horizontal",
                "Boundary",
                "Ratio",
                "OK",
                "Notes",
            ],
            [
                [
                    w.level,
                    w.mark,
                    w.t,
                    w.Lw,
                    w.Pu,
                    w.Mu,
                    w.Vu,
                    w.rho_v * 100,
                    w.rho_h * 100,
                    w.curtains,
                    w.vertical,
                    w.horizontal,
                    w.boundary,
                    w.utilisation,
                    yes[bool(w.ok)],
                    "; ".join(w.notes),
                ]
                for w in rep.walls
            ],
            bad_col=14,
        )
        self._fill(
            "Combined footings",
            ["Columns", "P kN", "L m", "B m", "D m", "q kN/m²", "Top", "Bottom", "Transverse", "OK", "Notes"],
            [
                [
                    " + ".join(c.marks),
                    c.P_service,
                    c.L,
                    c.B,
                    c.D,
                    c.q,
                    c.top,
                    c.bottom,
                    "; ".join(c.transverse),
                    yes[bool(c.ok)],
                    "; ".join(c.notes),
                ]
                for c in rep.combined_footings
            ],
            bad_col=9,
        )
        self._fill(
            "IS 13920",
            ["Level", "Member", "Type", "Clause / check", "OK", "Detail"],
            [
                [d.level, d.mark, d.kind, name, yes[bool(ok)], detail]
                for d in rep.ductile
                for name, ok, detail in d.checks
            ],
            bad_col=4,
        )
        irr = list(rep.irregularities) or list(getattr(fa, "irregularities", None) or [])
        self._fill(
            "Irregularity",
            ["Table", "Irregularity", "Clause", "Status", "Detail"],
            [[i.table, i.name, i.clause, i.status, i.detail] for i in irr],
        )
        self._fill_modal(rep, fa, fm)
        boq = rep.boq or {}
        rows = [
            [group, k, v.get("concrete", 0.0), v.get("steel", 0.0), v.get("formwork", 0.0), v.get("cost", 0.0)]
            for group, sec in (("Floor", boq.get("by_level", {})), ("Member type", boq.get("by_type", {})))
            for k, v in sec.items()
        ]
        if rows:
            rows.append(
                [
                    "Total",
                    "Building",
                    boq.get("total_concrete", 0.0),
                    boq.get("total_steel", 0.0),
                    boq.get("formwork", 0.0),
                    boq.get("cost", 0.0),
                ]
            )
        self._fill("BOQ by floor", ["Group", "Level / type", "Concrete m³", "Steel kg", "Formwork m²", "Cost ₹"], rows)

    def _fill_modal(self, rep, fa, fm):
        modal = rep.modal or getattr(fa, "modal", None)
        rs = getattr(fa, "rs", None) or getattr(fm, "rs", None) or {}
        headers = [
            "Mode / case",
            "T s",
            "f Hz",
            "UX %",
            "UY %",
            "RZ %",
            "ΣUX %",
            "ΣUY %",
            "Modes used",
            "Vb dynamic kN",
            "Vb static kN",
            "Scale",
        ]
        method = {"static": "equivalent static (cl 7.6)", "response spectrum": "response spectrum (cl 7.7)"}
        rows: list[list] = [["Method: " + method.get(rep.seismic_method, str(rep.seismic_method))] + [""] * 11]
        if modal is not None:
            cx, cy = modal.cumulative("x"), modal.cumulative("y")
            for k, md in enumerate(modal.modes):
                rows.append(
                    [
                        f"Mode {md.number}",
                        md.period,
                        1 / md.period if md.period > 0 else 0.0,
                        md.mass_x * 100,
                        md.mass_y * 100,
                        md.mass_rz * 100,
                        cx[k] * 100,
                        cy[k] * 100,
                    ]
                    + [""] * 4
                )
        for case, r in rs.items():
            rows.append([case] + [""] * 7 + [r.modes_used, r.vb_dynamic, r.vb_static, r.scale])
        self._fill("Modal / RS", headers, rows)

    def _issue_clicked(self, row, _col):
        if 0 <= row < len(self._issues):
            self.issueActivated.emit(self._issues[row])
