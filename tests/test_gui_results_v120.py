"""Off-screen tests of the grouped results navigator (tree + table, filter, copy, Excel export)."""

import copy
import os

import pytest

pytest.importorskip("PySide6")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from planwin_ai import units  # noqa: E402
from planwin_ai.ai.templates import build_template  # noqa: E402
from planwin_ai.units import T_TO_KN  # noqa: E402


@pytest.fixture(scope="module")
def app():
    from PySide6.QtWidgets import QApplication

    return QApplication.instance() or QApplication([])


@pytest.fixture(scope="module")
def designed():
    from planwin_ai.ai.actions import Session, execute

    s = Session(build_template("bungalow"))
    execute(s, [{"action": "design"}])
    return s.project, s.last["fm"], s.last["fa"], s.last["rep"]


class FakeMain:
    """The parts of MainWindow the results panel uses."""

    def __init__(self, project, fm=None, fa=None, rep=None, plan="Typical"):
        self.project, self.fm, self.fa, self.rep = project, fm, fa, rep
        self.current_plan_name = plan
        self.messages: list[str] = []

    def current_plan(self):
        return self.project.plan(self.current_plan_name)

    def plan_result(self, name):
        from planwin_ai.core.plan_engine import PlanEngine

        return PlanEngine(self.project.plan(name)).run()

    def frame_model(self):
        return self.fm

    def frame_analysis(self):
        return self.fa

    def design_report(self):
        return self.rep

    def statusBar(self):  # noqa: N802 - Qt name
        main = self

        class Bar:
            def showMessage(self, msg, _ms=0):  # noqa: N802
                main.messages.append(msg)

        return Bar()


@pytest.fixture
def panel(app, designed):
    from planwin_ai.gui.results_panel import ResultsPanel

    p, fm, fa, rep = designed
    rp = ResultsPanel(FakeMain(p, fm, fa, rep))
    rp.refresh()
    yield rp
    rp.deleteLater()


def test_tables_tree_and_show_tab(panel):
    from PySide6.QtCore import Qt

    from planwin_ai.gui.panels import RESULT_TABS, ResultsPanel
    from planwin_ai.gui.results_panel import RESULT_GROUPS, TABLE_INFO

    assert isinstance(panel, ResultsPanel) and list(panel.tables) == list(RESULT_TABS)
    assert panel.nav.objectName() == "ResultsNav" and panel.filter.objectName() == "ResultsFilter"
    grouped = [n for _g, names in RESULT_GROUPS for n in names]
    assert sorted(grouped) == sorted(RESULT_TABS) and set(TABLE_INFO) == set(RESULT_TABS)
    groups = [panel.nav.topLevelItem(i) for i in range(panel.nav.topLevelItemCount())]
    assert [g.text(0) for g in groups] == ["Plan", "Analysis", "Design", "Code checks", "Quantities"]
    assert all(g.isExpanded() and not (g.flags() & Qt.ItemIsSelectable) for g in groups)
    for name in ("IS 13920", "Lateral", "BOQ by floor", "Issues"):
        panel.show_tab(name)
        assert panel.stack.currentWidget() is panel.tables[name]
        assert panel.nav.currentItem().text(0) == name and panel.nav.currentItem().isSelected()
        assert panel.current_name() == name and panel.tabText(panel.currentIndex()) == name
    assert "Columns" in panel._items["Columns"].toolTip(0) and "biaxial" in panel._items["Columns"].toolTip(0)
    # clicking a group toggles it and keeps the current table
    panel.show_tab("Beams")
    design = groups[2]
    panel._nav_clicked(design, 0)
    assert not design.isExpanded() and panel.current_name() == "Beams"
    panel._nav_clicked(design, 0)
    assert design.isExpanded()


def test_up_down_keys_skip_groups(panel):
    from PySide6.QtCore import QEvent, Qt
    from PySide6.QtGui import QKeyEvent
    from PySide6.QtWidgets import QApplication

    def key(k):
        QApplication.sendEvent(panel.nav, QKeyEvent(QEvent.KeyPress, k, Qt.NoModifier))

    panel.show_tab("Beam loads")
    key(Qt.Key_Down)  # over the "Analysis" group row
    assert panel.current_name() == "Lateral" and panel.stack.currentWidget() is panel.tables["Lateral"]
    key(Qt.Key_Up)
    assert panel.current_name() == "Beam loads"
    panel.show_tab("Issues")
    key(Qt.Key_Up)  # nothing above: stays
    assert panel.current_name() == "Issues"


def test_badges_count_failures_and_issues(app, designed):
    from PySide6.QtCore import Qt

    from planwin_ai.gui import theme
    from planwin_ai.gui.results_panel import ResultsPanel

    p, fm, fa, rep = designed
    bad = copy.deepcopy(rep)
    bad.columns[0].ok = False
    bad.columns[1].ok = False
    rp = ResultsPanel(FakeMain(p, fm, fa, bad))
    rp.refresh()
    item = rp._items["Columns"]
    assert item.text(1) == "✖ 2" and item.foreground(1).color().name().upper() == "#DC2626"
    beams = rp._items["Beams"]
    assert all(b.ok for b in rep.beams) and beams.text(1) == "✓"
    n = rp.tables["Issues"].rowCount()
    assert rp._items["Issues"].text(1) == (f"({n})" if n else "")
    rp._fill("Issues", ["Where", "Level", "Type", "Message"], [["P", "error", "x", "m"], ["P", "warning", "y", "n"]], 1)
    assert rp._items["Issues"].text(1) == "(2)"
    assert rp._items["Lateral"].text(1) == "" and rp._items["Column loads"].text(1) == ""
    # invalidated results: design tables empty, badges gone, item greyed but still selectable
    rp.main.rep = rp.main.fa = rp.main.fm = None
    rp.refresh()
    for name in ("Columns", "Walls", "Combined footings", "IS 13920", "Modal / RS", "BOQ by floor", "Lateral"):
        assert rp.tables[name].rowCount() == 0, name
    muted = theme.PALETTES["light"]["muted"].lower()
    assert item.text(1) == "" and item.foreground(0).color().name() == muted
    assert item.flags() & Qt.ItemIsSelectable and rp._items["Column loads"].foreground(0).color().name() != muted
    rp.show_tab("Columns")
    assert rp.area.currentWidget() is rp.empty and "Design all (F7)" in rp.empty.text()
    rp.show_tab("Lateral")
    assert "Analyse frame (F6)" in rp.empty.text()
    rp.show_tab("Column loads")
    assert rp.area.currentWidget() is rp.stack


def test_filter_and_table_rows(panel):
    from planwin_ai.gui.results_panel import table_rows

    panel.show_tab("Columns")
    t = panel.tables["Columns"]
    total = t.rowCount()
    assert panel.count_lbl.text() == f"{total} rows"
    mark = t.item(0, 1).text()
    panel.filter.setText(mark.lower())  # case-insensitive, any column
    shown = [r for r in range(total) if not t.isRowHidden(r)]
    assert 0 < len(shown) < total
    assert all(any(mark.lower() in t.item(r, c).text().lower() for c in range(t.columnCount())) for r in shown)
    assert panel.count_lbl.text() == f"{len(shown)} of {total} rows"
    headers, rows = table_rows(t)
    assert headers[:2] == ["Level", "Col"] and len(rows) == len(shown)
    assert isinstance(rows[0][4], float) and isinstance(rows[0][11], str)  # numbers stay numbers
    assert len(table_rows(t, visible_only=False)[1]) == total
    panel.filter.setText("no such text zz")
    assert panel.count_lbl.text() == f"0 of {total} rows"
    panel.filter.clear()
    assert panel.count_lbl.text() == f"{total} rows" and not any(t.isRowHidden(r) for r in range(total))


def test_copy_puts_tsv_on_clipboard(panel):
    from PySide6.QtWidgets import QApplication

    panel.show_tab("Footings")
    t = panel.tables["Footings"]
    text = panel.copy_rows()
    lines = QApplication.clipboard().text().splitlines()
    assert text.splitlines() == lines and len(lines) == t.rowCount() + 1
    assert lines[0].split("\t")[:2] == ["Col", "P kN"]
    t.clearSelection()
    t.selectRow(1)
    lines = panel.copy_rows().splitlines()
    assert len(lines) == 2 and lines[1].split("\t")[0] == t.item(1, 0).text()
    assert panel.main.messages and "Copied 1 row" in panel.main.messages[-1]


def test_write_tables_xlsx_round_trip(tmp_path, panel, monkeypatch):
    from openpyxl import load_workbook
    from PySide6.QtWidgets import QFileDialog

    from planwin_ai.gui.results_panel import sheet_title, table_rows, write_tables_xlsx

    path = str(tmp_path / "r.xlsx")
    data = {
        "Columns": (["Col", "Pu kN", "OK"], [["C1", 812.4, "YES"], ["C3", 1120, "NO"]]),
        "BOQ & cost / [x]: a very long table name indeed": (["Item"], [["Concrete"]]),
    }
    write_tables_xlsx(path, data)
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "Columns" and len(wb.sheetnames[1]) <= 31 and "/" not in wb.sheetnames[1]
    ws = wb["Columns"]
    assert [c.value for c in ws[1]] == ["Col", "Pu kN", "OK"] and ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert ws["B2"].value == 812.4 and isinstance(ws["B3"].value, int) and ws["C3"].value == "NO"
    assert ws.column_dimensions["A"].width >= 4
    assert sheet_title("a:b", {"a_b"}) == "a_b (2)"
    # the panel exports the visible rows of the current table, and all tables with the menu entry
    panel.show_tab("Columns")
    panel.filter.setText(panel.tables["Columns"].item(0, 1).text())
    out = str(tmp_path / "cols.xlsx")
    seen = {}

    def ask(_parent, _title, default, _filter):
        seen["default"] = default
        return out, ""

    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(ask))
    assert panel.export_current() == out
    assert os.path.basename(seen["default"]).endswith("_Columns.xlsx")
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert len(rows) == 1 + len(table_rows(panel.tables["Columns"])[1])
    panel.filter.clear()
    assert panel.export_all() == out
    wb = load_workbook(out)
    assert "Columns" in wb.sheetnames and "IS 13920" in wb.sheetnames and "Lateral" in wb.sheetnames
    assert any("Exported" in m for m in panel.main.messages)


def test_mks_converts_kn_headers(app, designed):
    from planwin_ai.gui.results_panel import ResultsPanel

    p, fm, fa, rep = designed
    rp = ResultsPanel(FakeMain(p, fm, fa, rep))
    units.set_system("SI")
    rp.refresh()
    pu_si = float(rp.tables["Columns"].item(0, 4).text())
    try:
        units.set_system("MKS")
        rp.refresh()
        cols = rp.tables["Columns"]
        assert cols.horizontalHeaderItem(4).text() == "Pu t"
        assert float(cols.item(0, 4).text()) == pytest.approx(pu_si / T_TO_KN, abs=2e-3)
    finally:
        units.set_system("SI")


def test_issue_double_click_emits_issue(panel):
    got = []
    panel.issueActivated.connect(got.append)
    panel._issues = ["first", "second"]
    panel._fill("Issues", ["Where", "Level", "Type", "Message"], [["a", "error", "k", "m"], ["b", "info", "k", "m"]], 1)
    assert not panel.tables["Issues"].isSortingEnabled()
    panel.tables["Issues"].cellDoubleClicked.emit(1, 0)
    assert got == ["second"]


def test_settings_remember_table_and_splitter(app, designed, tmp_path):
    from PySide6.QtCore import QSettings

    from planwin_ai.gui.results_panel import ResultsPanel

    p, fm, fa, rep = designed
    main = FakeMain(p, fm, fa, rep)
    main.settings = QSettings(str(tmp_path / "s.ini"), QSettings.IniFormat)
    rp = ResultsPanel(main)
    rp.show_tab("Drift")
    rp.splitter.setSizes([240, 700])
    rp._save_splitter()
    rp2 = ResultsPanel(main)
    assert rp2.current_name() == "Drift" and rp2.stack.currentWidget() is rp2.tables["Drift"]
    assert [int(x) for x in main.settings.value("results/splitter")][0] == rp.splitter.sizes()[0]
