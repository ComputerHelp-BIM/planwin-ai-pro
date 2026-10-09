"""Bar bending schedule (planwin_ai.io.bbs)."""

import os

import pytest

from planwin_ai.ai.templates import build_template
from planwin_ai.core.model import grade_fck
from planwin_ai.design.runner import run_full
from planwin_ai.io.bbs import (
    BarItem,
    build_bbs,
    closed_link_length,
    development_length,
    lap_length,
    laps_needed,
    unit_weight,
    write_bbs_excel,
)


@pytest.fixture(scope="module")
def bungalow():
    prj = build_template("bungalow")
    fm, _fa, rep = run_full(prj)
    return prj, fm, rep, build_bbs(prj, fm, rep)


def test_unit_weight_and_item_weight():
    assert unit_weight(16) == pytest.approx(1.58, abs=0.001)  # 16² / 162
    it = BarItem("Beam", "B1", "L1", "B1-L1-01", "Bottom main", "00", 16, 3, 2, 4.5, {"a": 4500})
    assert it.total_length == pytest.approx(27.0)
    assert it.weight == pytest.approx(27.0 * 256 / 162)


def test_development_and_lap_length():
    # IS 456 cl 26.2.1.1: M25 τbd = 1.4 × 1.6 (deformed); Fe 500 -> σs = 0.87 × 500
    ld = 16 * 0.87 * 500 / (4 * 1.4 * 1.6)
    assert development_length(16, 25, 500) == pytest.approx(ld)
    assert lap_length(16, 25, 500) == pytest.approx(max(ld, 30 * 16))
    assert lap_length(8, 20, 250) == pytest.approx(max(8 * 0.87 * 250 / (4 * 1.2 * 1.6), 30 * 8))
    assert laps_needed(11000, 600) == 0
    assert laps_needed(13000, 600) == 1


def test_closed_stirrup_cutting_length():
    # 230 x 450 beam, 25 mm cover, T8 closed stirrup with two 135° hooks
    b, D, c, d = 230, 450, 25, 8
    a, h = b - 2 * c, D - 2 * c  # 180 x 400 out-to-out
    hook = max(10 * d, 75)  # 80
    deductions = 3 * 2 * d + 2 * 3 * d  # three 90° bends (2d) + two 135° hooks (3d)
    expected = 2 * a + 2 * h + 2 * hook - deductions
    assert expected == 1224
    assert closed_link_length(a, h, d) == pytest.approx(expected)


def test_bungalow_stirrups_use_rule(bungalow):
    prj, _fm, rep, bbs = bungalow
    bd = next(b for b in rep.beams if abs(b.b - 0.23) < 1e-9 and abs(b.d - 0.45) < 1e-9 and b.links)
    it = next(i for i in bbs.items if i.member == bd.mark and i.description == "Stirrup" and i.level == bd.level)
    c = prj.design.beam_cover * 1000
    assert it.shape_code == "51"
    assert it.cutting_length == pytest.approx(closed_link_length(230 - 2 * c, 450 - 2 * c, bd.links.dia) / 1000, 1e-3)


def test_column_bar_length_is_storey_height_plus_lap(bungalow):
    prj, fm, rep, bbs = bungalow
    # an intermediate storey (column above and below): bars = storey height + lap
    above = {m.n1 for m in fm.members.values() if m.kind == "column"}
    cd = next(
        c for c in rep.columns if c.level_index > 1 and fm.members[c.member_id].n2 in above and c.main_bars is not None
    )
    fck = grade_fck(fm.members[cd.member_id].grade)
    lap = lap_length(cd.main_bars.dia, fck, prj.design.fy_main)
    it = next(i for i in bbs.items if i.member_type == "Column" and i.member == cd.mark and i.level == cd.level)
    assert it.shape_code == "00"
    assert it.dia == cd.main_bars.dia and it.count == cd.main_bars.count
    assert it.cutting_length == pytest.approx((cd.height * 1000 + lap) / 1000, abs=1e-3)


def test_bungalow_bbs_totals(bungalow):
    _prj, _fm, rep, bbs = bungalow
    total = bbs.total_weight
    assert total > 0
    assert all(it.weight > 0 and it.cutting_length > 0 for it in bbs.items)
    assert sum(bbs.by_dia().values()) == pytest.approx(total)
    assert sum(bbs.by_level().values()) == pytest.approx(total)
    assert sum(bbs.by_member_type().values()) == pytest.approx(total)
    assert set(bbs.by_member_type()) == {"Beam", "Column", "Footing", "Slab"}
    marks = [it.bar_mark for it in bbs.items]
    assert len(marks) == len(set(marks))
    boq = rep.boq["total_steel"]
    assert 0.4 * boq <= total <= 1.6 * boq
    assert bbs.steel_per_m3 > 0


def test_plan_beam_segments_are_aggregated():
    from planwin_ai.core.model import Beam, Column, Level, Plan, Project

    plan = Plan(name="P")
    plan.columns = [Column(mark=f"C{i + 1}", x=x, y=y) for i, (x, y) in enumerate([(0, 0), (4, 0), (8, 0)])]
    plan.columns += [Column(mark=f"C{i + 4}", x=x, y=4) for i, x in enumerate([0, 4, 8])]
    plan.beams = [
        Beam(mark="B1", x1=0, y1=0, x2=8, y2=0),  # continuous over C2 -> two frame segments
        Beam(mark="B2", x1=0, y1=4, x2=8, y2=4),
        Beam(mark="B3", x1=0, y1=0, x2=0, y2=4),
        Beam(mark="B4", x1=4, y1=0, x2=4, y2=4),
        Beam(mark="B5", x1=8, y1=0, x2=8, y2=4),
    ]
    prj = Project(plans=[plan], levels=[Level("L1", "P", 3.0)])
    fm, _fa, rep = run_full(prj)
    segs = [b for b in rep.beams if b.mark == "B1"]
    assert len(segs) == 2
    bbs = build_bbs(prj, fm, rep)
    b1 = [it for it in bbs.items if it.member == "B1"]
    bottom = [it for it in b1 if it.description == "Bottom main"]
    assert len(bottom) == 1  # one physical beam
    assert bottom[0].dims["b"] > 8000  # runs the summed span (+ into the end columns)
    assert any("interior support" in it.description for it in b1)
    assert sum(1 for it in b1 if it.description.startswith("Hanger")) == 2
    # IS 13920 applies (zone III): closer links over 2d from each support, design spacing between
    from planwin_ai.io.bbs import zoned_counts

    c = prj.design.beam_cover * 1000
    stirrups = sum(it.count for it in b1 if it.description.startswith("Stirrup"))
    expected = 0
    for sg in segs:
        z = 2 * (sg.d * 1000 - c - 25)
        expected += sum(n for _, n in zoned_counts(4000, sg.links, sg.links_end, z, z))
    assert stirrups == expected
    assert stirrups >= sum(int(-(-4000 // sg.links.spacing)) + 1 for sg in segs)
    marks = [it.bar_mark for it in bbs.items]
    assert len(marks) == len(set(marks))


def test_bbs_excel(bungalow, tmpdir_path):
    from openpyxl import load_workbook

    prj, _fm, _rep, bbs = bungalow
    path = write_bbs_excel(os.path.join(tmpdir_path, "bbs.xlsx"), bbs, prj, watermark="TRIAL VERSION")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Beams", "Columns", "Footings", "Slabs"]
    ws = wb["Beams"]
    assert ws.freeze_panes == "A2"
    head = [c.value for c in ws[1]]
    assert head[:3] == ["Member", "Level", "Bar mark"] and "Weight (kg)" in head
    values = [c.value for row in wb["Summary"].iter_rows() for c in row if c.value is not None]
    assert "TRIAL VERSION" in values
    assert any("Verify against approved structural drawings before cutting" in str(v) for v in values)
    n_beam_rows = sum(1 for it in bbs.items if it.member_type == "Beam")
    assert ws.cell(n_beam_rows + 2, 1).value == "TOTAL"
    assert ws.cell(n_beam_rows + 2, 13).value == pytest.approx(bbs.by_member_type()["Beam"], abs=0.05)
