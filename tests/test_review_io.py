"""File I/O and export review: STAAD / ETABS syntax, DXF, legacy import, project files, BBS and the registry.

Structural checks parse the generated files the way the target program would (joints before
members, load cases before combinations, quoted names balanced ...) instead of comparing text.
"""

import json
import os
import re

import ezdxf
import pytest

from planwin_ai.ai.templates import build_template
from planwin_ai.core.generator import GridSpec, grid_building
from planwin_ai.core.model import Wall
from planwin_ai.core.solver import MLoad, MPoint
from planwin_ai.design.runner import run_full
from planwin_ai.io import project_io
from planwin_ai.io.bbs import build_bbs
from planwin_ai.io.detail_dxf import build_detail_doc
from planwin_ai.io.dxf_io import export_frame_dxf, export_plan_dxf, import_dxf
from planwin_ai.io.etabs import write_etabs
from planwin_ai.io.legacy_plw import T_TO_KN, LegacyFormatError, read_plw
from planwin_ai.io.staad import write_staad
from planwin_ai.services import exports

SAMPLES = os.path.join(os.path.dirname(__file__), "..", "planwin_ai", "data", "legacy_samples")


# ----------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="module")
def walled():
    """Bungalow with a shear wall on every plan (wide column + rigid links in the frame)."""
    p = build_template("bungalow")
    for pl in p.plans:
        pl.walls.append(Wall(mark="W1", x1=0.5, y1=0.0, x2=3.5, y2=0.0, thickness=0.2))
    fm, fa, rep = run_full(p)
    assert rep.walls
    return p, fm, fa, rep


@pytest.fixture(scope="module")
def combined():
    """Every column sits on a combined footing (no isolated footings at all)."""
    p = grid_building(GridSpec(bays_x=[2.0], bays_y=[6.0], upper_floors=4, city="Pune"))
    p.design.sbc = 120.0
    fm, fa, rep = run_full(p)
    assert rep.combined_footings and not rep.footings
    return p, fm, fa, rep


def _ctx(prj, fm, fa, rep, out_dir, params=None):
    return exports.ExportContext(prj, str(out_dir), "TRIAL", params or {}, lambda: (fm, fa), lambda: rep, lambda: {})


# ----------------------------------------------------------------------------- STAAD
def _staad_logical_lines(text: str) -> list[str]:
    out, cur = [], ""
    for ln in text.splitlines():
        if ln.endswith(" -"):  # continuation
            cur += ln[:-2] + " "
        else:
            out.append(cur + ln)
            cur = ""
    return out


def _expand(tokens: list[str]) -> list[int]:
    out, i = [], 0
    while i < len(tokens):
        if i + 2 < len(tokens) and tokens[i + 1] == "TO":
            out += range(int(tokens[i]), int(tokens[i + 2]) + 1)
            i += 3
        else:
            out.append(int(tokens[i]))
            i += 1
    return out


def parse_staad(text: str) -> dict:
    joints, members, supports, slaves, loads, combos = {}, {}, set(), [], [], {}
    load_totals: dict[int, list[float]] = {}
    sec = None
    for ln in _staad_logical_lines(text):
        u = ln.upper().strip()
        if u in ("JOINT COORDINATES", "MEMBER INCIDENCES", "SUPPORTS", "MEMBER LOAD", "JOINT LOAD"):
            sec = u
            continue
        if u.startswith("LOAD COMB"):
            sec, cur = "COMB", int(u.split()[2])
            combos[cur] = []
            continue
        if re.match(r"LOAD \d+", u):
            sec, cur = None, int(u.split()[1])
            loads.append(cur)
            load_totals[cur] = [0.0, 0.0, 0.0]
            continue
        if u.startswith("SLAVE"):
            t = u.split()
            slaves.append((int(t[t.index("MASTER") + 1]), _expand(t[t.index("JOINT") + 1 :])))
            continue
        if not u or u.startswith("*"):
            continue
        if not u[0].isdigit():  # any other command ends the data block
            sec = None
            continue
        t = u.rstrip(";").split()
        if sec == "JOINT COORDINATES":
            assert int(t[0]) not in joints, f"duplicate joint {t[0]}"
            joints[int(t[0])] = tuple(map(float, t[1:4]))
        elif sec == "MEMBER INCIDENCES":
            assert int(t[0]) not in members, f"duplicate member {t[0]}"
            members[int(t[0])] = (int(t[1]), int(t[2]))
        elif sec == "SUPPORTS":
            supports |= set(_expand(t[:-1]))
        elif sec == "COMB":
            combos[cur] += [int(t[k]) for k in range(0, len(t), 2)]
        elif sec == "MEMBER LOAD":
            assert int(t[0]) in members, ln
            tot = load_totals[loads[-1]]
            if t[1] == "UNI":
                tot[2] += float(t[3]) * (float(t[5]) - float(t[4]))
            elif t[1] == "TRAP":
                tot[2] += (float(t[3]) + float(t[4])) / 2 * (float(t[6]) - float(t[5]))
            elif t[1] == "CON":
                tot[2] += float(t[3])
        elif sec == "JOINT LOAD":
            assert int(t[0]) in joints, ln
            d = dict(zip(t[1::2], map(float, t[2::2])))
            tot = load_totals[loads[-1]]
            # STAAD (X, Y up, Z) -> plan (x, y = -Z, z = Y)
            tot[0] += d.get("FX", 0.0)
            tot[1] -= d.get("FZ", 0.0)
            tot[2] += d.get("FY", 0.0)
    return {
        "joints": joints,
        "members": members,
        "supports": supports,
        "slaves": slaves,
        "loads": loads,
        "combos": combos,
        "totals": load_totals,
    }


def _model_totals(fm) -> dict[str, tuple[float, float, float]]:
    nodal, _ = fm.export_model()
    out = {}
    for c in fm.cases():
        fx = fy = fz = 0.0
        for m in fm.members.values():
            for ld in m.loads.get(c, []):
                if isinstance(ld, MLoad):
                    fz += (ld.w1[2] + ld.w2[2]) / 2 * (ld.b - ld.a)
                elif isinstance(ld, MPoint):
                    fz += ld.P[2]
        for v in nodal.get(c, {}).values():
            fx, fy, fz = fx + v[0], fy + v[1], fz + v[2]
        out[c] = (fx, fy, fz)
    return out


def test_staad_file_is_structurally_valid(walled, tmp_path):
    p, fm, _fa, _rep = walled
    text = open(write_staad(fm, str(tmp_path / "w.std")), encoding="ascii").read()
    s = parse_staad(text)
    assert list(s["joints"]) == list(range(1, len(s["joints"]) + 1))
    assert list(s["members"]) == list(range(1, len(s["members"]) + 1))
    used = {n for m in s["members"].values() for n in m}
    assert used == set(s["joints"])  # no orphan joints
    assert all(a != b for a, b in s["members"].values())
    assert s["supports"] and s["supports"] <= set(s["joints"])
    for master, sl in s["slaves"]:
        assert master not in sl and {master, *sl} <= set(s["joints"])
        assert not ({master, *sl} & s["supports"])
        assert len({s["joints"][j][1] for j in (master, *sl)}) == 1  # one floor (STAAD Y up)
    assert s["loads"] == list(range(1, len(fm.cases()) + 1))
    assert s["combos"] and all(refs and set(refs) <= set(s["loads"]) for refs in s["combos"].values())
    # every load the model carries reaches the file (no joint load dropped, signs mapped x, z, -y)
    want = _model_totals(fm)
    for k, case in enumerate(fm.cases(), start=1):
        assert s["totals"][k] == pytest.approx(want[case], abs=0.05 + 1e-4 * max(map(abs, want[case])))


def test_staad_free_text_cannot_break_the_input(walled, tmp_path):
    """';' separates STAAD commands, so a project/client name must not contain one; lines stay within 79."""
    p, fm, _fa, _rep = walled
    old = (p.name, p.client, p.engineer)
    p.name, p.client, p.engineer = "Block A; Wing B", "Shah; Sons", "R. Rao; PE"
    try:
        text = open(write_staad(fm, str(tmp_path / "t.std")), encoding="ascii").read()
    finally:
        p.name, p.client, p.engineer = old
    head = text.split("END JOB INFORMATION")[0]
    assert ";" not in head and "JOB NAME Block A" in head
    assert all(len(ln) <= 79 for ln in text.splitlines())


# ----------------------------------------------------------------------------- ETABS
def _q(line: str) -> list[str]:
    return re.findall(r'"([^"]*)"', line)


def test_etabs_file_is_structurally_valid(walled, tmp_path):
    p, fm, _fa, _rep = walled
    lines = open(write_etabs(fm, str(tmp_path / "w.e2k")), encoding="utf-8").read().splitlines()
    assert all(ln.count('"') % 2 == 0 for ln in lines)
    stories, points, frames, sections, patterns, cases = set(), set(), {}, set(), set(), set()
    for ln in (x.strip() for x in lines):
        q = _q(ln)
        if ln.startswith("STORY "):
            assert q[0] not in stories
            stories.add(q[0])
        elif ln.startswith("POINT "):
            assert q[0] not in points
            points.add(q[0])
        elif ln.startswith("FRAMESECTION ") and "MATERIAL" in ln:
            sections.add(q[0])
        elif ln.startswith("LINE "):
            assert q[0] not in frames and {q[1], q[2]} <= points
            frames[q[0]] = ln.split()[2]
        elif ln.startswith("POINTASSIGN "):
            assert q[0] in points and q[1] in stories
        elif ln.startswith("LINEASSIGN "):
            assert q[0] in frames and q[1] in stories and q[2] in sections
        elif ln.startswith("LOADPATTERN "):
            patterns.add(q[0])
        elif ln.startswith(("POINTLOAD ", "LINELOAD ")):
            assert q[1] in stories and re.search(r'LC "([^"]*)"', ln).group(1) in patterns
            assert (q[0] in points) if ln.startswith("POINTLOAD") else (q[0] in frames)
        elif ln.startswith("LOADCASE ") and "LOADPAT" in ln:
            assert q[1] in patterns
            cases.add(q[0])
        elif ln.startswith("COMBO ") and "LOADCASE" in ln:
            assert q[1] in cases
    assert len(frames) == len(fm.members) and "Base" in stories
    assert sum(k == "COLUMN" for k in frames.values()) == sum(m.kind in ("column", "wall") for m in fm.members.values())


def test_etabs_title_with_quotes_stays_one_string(walled, tmp_path):
    p, fm, _fa, _rep = walled
    old = p.name
    p.name = 'Tower "A"'
    try:
        lines = open(write_etabs(fm, str(tmp_path / "q.e2k")), encoding="utf-8").read().splitlines()
    finally:
        p.name = old
    title = next(ln for ln in lines if ln.strip().startswith("TITLE2"))
    assert title.count('"') == 2 and "Tower" in title


# ----------------------------------------------------------------------------- DXF
def test_plan_dxf_round_trip_and_audit(tmp_path):
    p = build_template("bungalow")
    for plan in p.plans:
        path = export_plan_dxf(plan, str(tmp_path / f"{plan.name}.dxf"))
        assert not ezdxf.readfile(path).audit().has_errors
        q, notes = import_dxf(path)
        assert not notes
        assert len(q.slabs) == len(plan.slabs) and len(q.beams) == len(plan.beams)

        def key(c):
            return round(c.x, 3), round(c.y, 3), round(c.b, 3), round(c.d, 3), round(c.angle % 180, 1)

        assert sorted(map(key, q.columns)) == sorted(map(key, plan.columns))
        area = sorted(round(abs(s.area), 2) for s in plan.slabs)
        assert sorted(round(abs(s.area), 2) for s in q.slabs) == area


def test_plan_dxf_draws_shear_walls(walled, tmp_path):
    p, *_ = walled
    plan = p.plans[1]
    doc = ezdxf.readfile(export_plan_dxf(plan, str(tmp_path / "w.dxf")))
    assert not doc.audit().has_errors
    walls = [e for e in doc.modelspace() if e.dxf.layer == "WALL" and e.dxftype() == "LWPOLYLINE"]
    assert len(walls) == len(plan.walls)
    xs = [x for e in walls for x, *_ in e.get_points()]
    assert min(xs) == pytest.approx(0.5) and max(xs) == pytest.approx(3.5)
    assert any(e.dxf.layer == "WALL_TEXT" and e.dxf.text == "W1" for e in doc.modelspace().query("TEXT"))
    q, _ = import_dxf(str(tmp_path / "w.dxf"))  # walls are not read back as slabs / columns / beams
    assert len(q.slabs) == len(plan.slabs) and len(q.columns) == len(plan.columns)


def test_frame_dxf_puts_vertical_members_on_the_column_layer(walled, tmp_path):
    _p, fm, _fa, _rep = walled
    doc = ezdxf.readfile(export_frame_dxf(fm, str(tmp_path / "f.dxf")))
    assert not doc.audit().has_errors
    lines = list(doc.modelspace().query("LINE"))
    assert len(lines) == len(fm.members)
    for e in lines:
        vertical = abs(e.dxf.start.z - e.dxf.end.z) > 1e-6
        assert e.dxf.layer == ("FRAME_COLUMN" if vertical else "FRAME_BEAM")


def test_detail_drawings_list_combined_footings(combined):
    p, fm, _fa, rep = combined
    doc, _info = build_detail_doc(p, fm, rep)
    assert not doc.audit().has_errors
    texts = [e.dxf.text for e in doc.modelspace().query("TEXT")]
    assert "No footing designs available." not in texts
    for c in rep.combined_footings:
        assert any(" + ".join(c.marks) in t for t in texts)


def test_pdf_report_title_uses_only_standard_font_glyphs(walled, tmp_path):
    """Helvetica has no ₹ / √ glyph: the report title is transliterated like every table cell."""
    pypdf = pytest.importorskip("pypdf")
    from planwin_ai.io.pdf_report import write_pdf

    p, fm, fa, rep = walled
    old = p.name
    p.name = "Villa <A&B> ₹ 50 lakh"
    try:
        path = write_pdf(str(tmp_path / "r.pdf"), p, {}, fm, rep, "TRIAL", fa=fa)
    finally:
        p.name = old
    first = pypdf.PdfReader(path).pages[0].extract_text()
    assert "Villa <A&B> INR 50 lakh" in first


# ----------------------------------------------------------------------------- legacy .plw
def test_legacy_loads_are_converted_from_tonnes():
    path = os.path.join(SAMPLES, "Box.plw")
    plan, rep = read_plw(path)
    # v1.x: first slab record "id, n, angle, mark, total load (t/m²), thickness, grade, direction, live (t/m²) ..."
    with open(path, encoding="latin-1") as f:
        toks = [t.strip().strip('"') for ln in f.readlines()[1:] for t in ln.split(",")]
    total_t, live_t = float(toks[12]), float(toks[16])
    s = plan.slabs[0]
    assert s.mark == toks[11]
    assert s.dead + s.live == pytest.approx(total_t * T_TO_KN, abs=1e-3)
    assert s.live == pytest.approx(live_t * T_TO_KN, abs=1e-3)
    assert T_TO_KN == pytest.approx(9.80665)


@pytest.mark.parametrize("body", ["\n1\n3,0,0\n", "\n1\n1,0,0\nx\nS1,4,0.15\n"])
def test_truncated_v5_file_raises_format_error(tmp_path, body):
    path = tmp_path / "v5.plw"
    path.write_text("Ver 5.7.1#01/01/10 10:00:00#A#1" + body, encoding="latin-1")
    with pytest.raises(LegacyFormatError):
        read_plw(str(path))


# ----------------------------------------------------------------------------- project files
def _write(path, data: bytes) -> str:
    with open(path, "wb") as f:
        f.write(data)
    return str(path)


def test_project_with_utf8_bom_opens(tmp_path):
    p = build_template("bungalow")
    text = project_io.project_to_json(p)
    q = project_io.load_project(_write(tmp_path / "bom.pwai", b"\xef\xbb\xbf" + text.encode("utf-8")))
    assert q.name == p.name and len(q.plans) == len(p.plans)


@pytest.mark.parametrize(
    "data",
    [
        b"\x89PNG\r\n\x1a\n\x00\xff",  # not text at all
        b'{"plans": "Ground"}',
        b'{"plans": [1]}',
        b'{"plans": [], "levels": [1]}',
        b'{"plans": [{"name": "A", "slabs": [5]}]}',
        b'{"plans": [], "seismic": "zone 3"}',
    ],
)
def test_corrupt_project_raises_format_error(tmp_path, data):
    with pytest.raises(project_io.ProjectFormatError):
        project_io.load_project(_write(tmp_path / "bad.pwai", data))


def test_newer_schema_message_is_kept(tmp_path):
    path = _write(tmp_path / "new.pwai", json.dumps({"plans": [], "schema": 999}).encode())
    with pytest.raises(project_io.ProjectFormatError, match="newer"):
        project_io.load_project(path)


def test_save_is_flushed_to_disk_before_it_replaces_the_old_file(tmp_path, monkeypatch):
    calls = []
    real_fsync, real_replace = os.fsync, os.replace
    monkeypatch.setattr(project_io.os, "fsync", lambda fd: (calls.append("fsync"), real_fsync(fd))[1])
    monkeypatch.setattr(project_io.os, "replace", lambda a, b: (calls.append("replace"), real_replace(a, b))[1])
    path = project_io.save_project(build_template("bungalow"), str(tmp_path / "x"))
    assert calls[:2] == ["fsync", "replace"]
    assert path.endswith(".pwai") and project_io.load_project(path).plans
    assert [f for f in os.listdir(tmp_path)] == ["x.pwai"]  # no temp file left behind


# ----------------------------------------------------------------------------- BBS
def test_bbs_flags_combined_footings_and_anchors_their_columns(combined):
    p, fm, _fa, rep = combined
    bbs = build_bbs(p, fm, rep)
    for c in rep.combined_footings:
        assert any("Combined footing" in w and " + ".join(c.marks) in w for w in bbs.warnings)
    on_combined = {m for c in rep.combined_footings for m in c.marks}
    bottom = {
        (cd.mark, cd.level) for cd in rep.columns if fm.nodes[fm.members[cd.member_id].n1].support
    }  # fmt: skip
    assert {m for m, _ in bottom} == on_combined
    starters = [it for it in bbs.items if it.description.startswith("Main bar") and (it.member, it.level) in bottom]
    assert len(starters) == len(bottom) and all("foot in footing" in it.description for it in starters)
    for it in bbs.items:  # totals by diameter are the sum of the rows, weight d²/162 per m
        assert it.unit_weight == pytest.approx(it.dia**2 / 162)
    assert sum(bbs.by_dia().values()) == pytest.approx(sum(it.weight for it in bbs.items))


def test_bbs_flags_shear_walls(walled):
    p, fm, _fa, rep = walled
    bbs = build_bbs(p, fm, rep)
    assert any(w.startswith("Wall W1") and "not scheduled" in w for w in bbs.warnings)


# ----------------------------------------------------------------------------- export registry
@pytest.mark.parametrize(
    "name,expected",
    [
        ("a" * 119 + ". tower", "a" * 119),  # no trailing dot after truncation (Windows drops it)
        ("LPT1", "_LPT1"),
        ("nul.txt", "_nul.txt"),
        ("", "project"),
        ("मुंबई टॉवर", "मुंबई_टॉवर"),
    ],
)
def test_safe_filename_edge_cases(name, expected):
    out = exports.safe_filename(name)
    assert out == expected
    assert not out.endswith((".", " ")) and len(out) <= 120


def test_export_plan_parameter_must_name_a_plan(walled, tmp_path):
    p, fm, fa, rep = walled
    ctx = _ctx(p, fm, fa, rep, tmp_path, {"plan": "typical"})
    assert ctx.plan() is p.plan("Typical")  # case does not matter
    assert exports.get("dxf").default_path(ctx).endswith("_Typical_2DPLAN.dxf")
    ctx = _ctx(p, fm, fa, rep, tmp_path, {"plan": "Basement"})
    with pytest.raises(ValueError, match="Basement"):
        exports.run(exports.get("dxf"), ctx)
    assert not os.listdir(tmp_path)


def test_every_export_runs_on_walls_and_combined_footings(walled, combined, tmp_path):
    names = [k for f in exports.formats() for k in (f.key, *f.aliases)]
    assert len(names) == len(set(names))
    for tag, (p, fm, fa, rep) in (("w", walled), ("c", combined)):
        params = {}
        if tag == "w":  # keep the calculation-sheet set small: one beam, one column, one footing, one slab
            params = {
                "member_ids": [rep.beams[0].member_id, rep.columns[0].member_id],
                "footing_marks": [rep.footings[0].mark],
                "slabs": [(rep.slabs[0][0], rep.slabs[0][1].mark)],
            }
        ctx = _ctx(p, fm, fa, rep, tmp_path / tag, params)
        paths = {fmt.key: exports.run(fmt, ctx) for fmt in exports.formats()}
        assert len(set(paths.values())) == len(paths)  # no two formats overwrite each other
        for key, path in paths.items():
            assert os.path.getsize(path) > 200, key
            if path.endswith(".dxf"):
                assert not ezdxf.readfile(path).audit().has_errors, key


def test_names_starting_with_equals_stay_text_in_every_workbook(tmp_path):
    """openpyxl stores any '=…' string as a formula: a project or level named '=Tower' gave Excel
    a broken formula (a 'repair this file' prompt, or formula injection from imported text)."""
    from openpyxl import load_workbook

    from planwin_ai.ai.actions import Session, execute
    from planwin_ai.ai.templates import build_template

    prj = build_template("bungalow")
    prj.name = "=Tower"
    prj.client = "=HYPERLINK(\"http://x\",\"y\")"
    prj.levels[0].name = "=L0"
    prj.meta["revisions"] = []
    s = Session(prj, out_dir=str(tmp_path))
    for key in ("excel", "bbs", "boq", "schedules"):
        res = execute(s, [{"action": "export", "format": key, "path": str(tmp_path / f"{key}.xlsx")}])
        assert not res.errors, res.errors
        wb = load_workbook(tmp_path / f"{key}.xlsx")
        bad = [
            (ws.title, c.coordinate, c.value)
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for c in row
            if c.data_type == "f" and str(c.value).startswith(("=Tower", "=HYPERLINK", "=L0"))
        ]
        assert not bad, (key, bad[:5])
    # the BOQ keeps its real formulas
    boq = load_workbook(tmp_path / "boq.xlsx")
    assert any(c.data_type == "f" for ws in boq.worksheets for row in ws.iter_rows() for c in row)


def test_results_table_excel_keeps_equals_text(tmp_path):
    from openpyxl import load_workbook

    from planwin_ai.gui.results_panel import write_tables_xlsx

    p = tmp_path / "t.xlsx"
    write_tables_xlsx(str(p), {"Beams": (["Mark", "Note"], [["=B1", "=1+1"]])})
    ws = load_workbook(p).active
    assert ws["A2"].data_type == "s" and ws["A2"].value == "=B1"
